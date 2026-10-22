"""Photo → Excel with the AI table (merged cells, numbers, dates, a formula) in real Excel
(NEXVARA_RUN_OFFICE=1): the workbook is built through COM and checked with openpyxl."""

import datetime as dt
import json
import os

import pytest

pytestmark = [
    pytest.mark.office,
    pytest.mark.skipif(
        os.environ.get("NEXVARA_RUN_OFFICE") != "1",
        reason="set NEXVARA_RUN_OFFICE=1 to drive real Office",
    ),
]

AI_ANSWER = json.dumps(
    {
        "title": "Beton tarkibi",
        "rows": [
            ["Sana", "Material", "Miqdor", "Narx"],
            ["2026-10-22", "Sement", 350, 1250.5],
            ["2026-10-22", "Qum", 700, 80.25],
            ["Jami", "", "", ""],
        ],
        "merges": ["A4:C4"],
        "formulas": {"D4": "=SUM(D2:D3)"},
    }
)


def test_ai_table_into_visible_excel(tmp_path):
    import openpyxl

    from nexvara import photodoc

    data = photodoc.parse_ai_table(AI_ANSWER)
    image = tmp_path / "rasm.jpg"
    image.write_bytes(b"\xff\xd8 fake")
    res = photodoc.convert(data, image, "excel", keep_open=False, use_com=True)
    wb = openpyxl.load_workbook(res["path"])
    ws = wb.active
    assert [c.value for c in ws[1]] == ["Sana", "Material", "Miqdor", "Narx"]
    assert ws["A2"].value.date() == dt.date(2026, 10, 22)  # a real date, not text
    assert ws["C2"].value == 350 and ws["D3"].value == 80.25  # real numbers
    assert ws["D4"].value == "=SUM(D2:D3)"
    assert "A4:C4" in {str(r) for r in ws.merged_cells.ranges}
    assert ws["A1"].font.bold
