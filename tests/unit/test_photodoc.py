"""Photo → Excel/Word: finding the table in OCR text, numbers, the .xlsx, chat flows."""

from __future__ import annotations

import openpyxl

from nexvara import photodoc


def test_table_rows_from_pipes_tabs_and_gaps():
    t = "O‘quvchilar ro‘yxati\nIsm | Sinf | Ball\n---|---|---\nAli | 7-A | 85\nVali | 7-B"
    rows = photodoc.table_rows(t)
    assert rows == [["Ism", "Sinf", "Ball"], ["Ali", "7-A", "85"], ["Vali", "7-B", ""]]
    assert photodoc.table_rows("Mahsulot    Narx\nNon    3 500\nSut    12 000") == [
        ["Mahsulot", "Narx"], ["Non", "3 500"], ["Sut", "12 000"]]  # fmt: skip
    assert photodoc.table_rows("Bu oddiy matn.\nIkkinchi qator ham oddiy.") is None


def test_numbers_and_texts():
    assert photodoc.number("1 250 000") == 1250000
    assert photodoc.number("12,5") == 12.5
    assert photodoc.number("007") == "007"  # leading zero: a code, not a number
    assert photodoc.number("7-A") == "7-A"
    assert photodoc.number("12345678901234567") == "12345678901234567"


def test_excel_has_header_borders_and_numbers(tmp_path):
    out = photodoc.write_excel([["Ism", "Ball"], ["Ali", "85"]], tmp_path / "x.xlsx", "Ro‘yxat")
    ws = openpyxl.load_workbook(out).active
    assert ws["A1"].font.bold and ws["B2"].value == 85
    assert ws["A2"].border.bottom.style == "thin" and ws.freeze_panes == "A2"


def test_chat_photo_to_excel_offline(tmp_path, monkeypatch):
    monkeypatch.setenv("NEXVARA_DATA", str(tmp_path / "data"))
    from nexvara.app.api import Api

    photo = tmp_path / "baholar.jpg"
    photo.write_bytes(b"\xff\xd8 x")
    api = Api(overrides={"ocr": lambda p, lang: "Ism\tBall\nAli\t5\nVali\t4"})
    api.set_context_file(str(photo))
    r = api.chat("shu rasmni Excel qil")
    assert r["reply"].startswith("📊") and "2 qator" in r["reply"], r
    assert (tmp_path / "baholar_NEXVARA.xlsx").is_file() and photo.read_bytes() == b"\xff\xd8 x"
    assert api.get_context_file().endswith("baholar_NEXVARA.xlsx")


def test_chat_photo_to_word_and_ai_table(tmp_path, monkeypatch):
    monkeypatch.setenv("NEXVARA_DATA", str(tmp_path / "data"))
    from test_drawing import Draftsman, _jpeg_with_gps

    from nexvara.ai.keys import MemoryKeyStore
    from nexvara.ai.providers import Reply
    from nexvara.ai.router import AiRouter, Ledger
    from nexvara.app.api import Api
    from nexvara.appdata import Settings
    from nexvara.core.task import TaskRunner

    photo = tmp_path / "xat.jpg"
    photo.write_bytes(_jpeg_with_gps())
    got = []
    api = Api(
        TaskRunner(),
        overrides={
            "ocr": lambda p, lang: "Hurmatli direktor!\n\nAriza.",
            "photo_word": lambda t, p: got.append(t) or {"path": "x"},
        },
    )
    api.set_context_file(str(photo))
    r = api.chat("rasmdagi matnni Word hujjatga aylantir")
    assert api._runner.get(r["task"]["id"]).wait(5) and "direktor" in got[0]

    class Reader(Draftsman):
        def complete(self, model, system, messages, max_tokens, effort="medium"):
            self.calls.append(messages)
            return Reply("Ism | Ball\nAli | 5\nVali | 4", model, 1500, 40)

    pro = Reader()
    router = AiRouter(
        Settings(tmp_path / "s.json"),
        MemoryKeyStore({"gemini": "AIza-test-1234567"}),
        Ledger(tmp_path / "l.db"),
        factory=lambda n, k: pro,
    )
    router.settings.update({"ai": {"enabled": True, "provider": "gemini"}})
    api2 = Api(overrides={"ai_router": router})
    api2.set_context_file(str(photo))
    ask = api2.chat("AI bilan shu rasmdagi jadvalni Excel qil")
    assert ask.get("confirm") and "Excel" in ask["reply"] and pro.calls == []
    done = api2.confirm_ai(ask["confirm"], True)
    assert done["reply"].startswith("📊") and "AI" in done["reply"], done
