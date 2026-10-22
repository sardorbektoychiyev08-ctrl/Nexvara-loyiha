"""AI-first chat routing. All providers here are in-memory fakes: no paid call is possible."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from nexvara import photodoc
from nexvara.ai.keys import MemoryKeyStore
from nexvara.ai.providers import AiError, Provider, Reply
from nexvara.ai.router import AiRouter, Ledger
from nexvara.appdata import Settings


class FakeProvider(Provider):
    def __init__(self, replies=(), error: Exception | None = None):
        self.replies = list(replies)
        self.error = error
        self.calls: list[dict] = []

    def complete(self, model, system, messages, max_tokens, effort="medium"):
        self.calls.append({"system": system, "messages": messages, "model": model})
        if self.error:
            raise self.error
        text = self.replies.pop(0) if self.replies else '{"kind":"chat"}'
        return Reply(text, model, 100, 20)

    def list_models(self):
        return ["claude-haiku-5-5", "claude-sonnet-5-5", "claude-opus-5-5"]


def _api(tmp_path: Path, replies=(), error: Exception | None = None, enabled: bool = True):
    from nexvara.app.api import Api

    fake = FakeProvider(replies, error)
    router = AiRouter(
        Settings(tmp_path / "settings.json"),
        MemoryKeyStore({"anthropic": "sk-ant-fake-only"}),
        Ledger(tmp_path / "usage.db"),
        factory=lambda _name, _key: fake,
    )
    router.settings.update({"ai": {"enabled": enabled, "ai_first": True, "confirm_over_usd": 1.0}})
    api = Api(
        overrides={
            "ai_router": router,
            "open_app": lambda app: f"{app} ochildi",
            "ocr": lambda _p, _lang: "Sana | Summa\n22.10.2026 | 1250",
        }
    )
    return api, fake, router


@pytest.mark.parametrize(
    ("message", "tool_name", "args"),
    [
        ("noutbukimdagi exselni och", "app.open", {"app": "excel"}),
        (
            "qara men hali senga yoz demadim tozala yozganlaringni",
            "files.recycle",
            {"paths": ["@prev"]},
        ),
        (
            "jadvalni chiroyliroq qilib tuz",
            "photo.to_office",
            {"target": "excel", "request": "jadvalni chiroyliroq qilib tuz"},
        ),
        (
            "bugungi sanadagini exselga jadval qilib tushib ber",
            "photo.to_office",
            {"target": "excel", "request": "faqat bugungi sana"},
        ),
        (
            "hozir qilgan ishingni ko‘ryabsanmi",
            "photo.compare_table",
            {"request": "hozirgi natijani tekshir"},
        ),
        (
            "qara jadvalda xatoliklar bor ko‘rgin",
            "photo.compare_table",
            {"request": "xatolarni top"},
        ),
        ("Excel’ni och", "app.open", {"app": "excel"}),
        (
            "Hisobot.xlsx ni och, B ustun bo‘yicha sarala",
            "nexvara.command",
            {"text": "Hisobot.xlsx ni och, B ustun bo‘yicha sarala"},
        ),
        (
            "7-sinf fizika 20 ta test",
            "education.test",
            {"request": "7-sinf fizika 20 ta test"},
        ),
    ],
)
def test_planner_is_first_and_selects_expected_tool(tmp_path, message, tool_name, args):
    plan = {"kind": "tools", "calls": [{"tool": tool_name, "args": args}], "say": "tushundim"}
    api, fake, _router = _api(tmp_path, [json.dumps(plan, ensure_ascii=False)])
    photo = tmp_path / "jadval.jpg"
    photo.write_bytes(b"\xff\xd8 fake")
    result = tmp_path / "jadval_NEXVARA.xlsx"
    photodoc.write_excel([["Sana", "Summa"], ["22.10.2026", "1250"]], result)
    api.set_context_file(str(photo))
    api._ctx.result = str(result)

    out = api.chat(message)

    assert fake.calls, message
    assert tool_name in fake.calls[0]["system"]
    assert fake.calls[0]["messages"][0]["content"].rstrip().endswith(message)
    if tool_name == "photo.compare_table":
        assert out.get("confirm") and api._pending[out["confirm"]][0] == "compare_table"
    else:
        assert out.get("confirm")
        kind, calls = api._pending[out["confirm"]]
        assert kind == "tools" and calls[0].tool == tool_name


def test_small_talk_also_goes_to_planner_first(tmp_path):
    api, fake, _router = _api(tmp_path, ['{"kind":"chat"}', "Salom, ishlaringiz yaxshimi?"])
    out = api.chat("salom")
    assert out["reply"].startswith("Salom")
    assert len(fake.calls) == 2
    assert "ASBOBLAR" in fake.calls[0]["system"]


def test_ai_disabled_keeps_old_handler_silently(tmp_path):
    """AI switched off by the owner: the old chain, no «AI hozir ishlamayapti» line."""
    api, fake, _router = _api(tmp_path, enabled=False)
    out = api.chat("Excel’ni och")
    assert out["reply"] == "Excel ochildi."
    assert fake.calls == [] and "ai_fallback" not in out


@pytest.mark.parametrize(
    "error",
    [
        AiError("Provayder serveri hozir band (503).", retryable=True),
        AiError("Gemini bepul chegarasi tugadi.", retryable=True),
    ],
)
def test_provider_failure_falls_back_to_old_chain(tmp_path, error):
    api, fake, _router = _api(tmp_path, error=error)
    out = api.chat("Excel’ni och")
    assert fake.calls and "Excel" in out["reply"]
    assert out["reply"].endswith("AI hozir ishlamayapti, oddiy rejimda bajardim")


def test_ai_table_json_builds_typed_office_free_workbook(tmp_path):
    raw = json.dumps(
        {
            "title": "Bugungi savdo",
            "rows": [
                ["Sana", "Mahsulot", "Summa", ""],
                ["2026-10-22", "Daftar", 1250.5, ""],
                ["", "", "", ""],
            ],
            "merges": ["A2:A3", "noto‘g‘ri"],
            "formulas": {"C3": "=SUM(C2:C2)", "bad": "=1"},
        },
        ensure_ascii=False,
    )
    data = photodoc.parse_ai_table(raw)
    assert data.title == "Bugungi savdo"
    assert data.rows[1][2] == 1250.5
    assert data.rows[1][0].isoformat() == "2026-10-22"
    assert len(data.rows[0]) == 3 and data.merges == ["A2:A3"]
    assert data.formulas == {"C3": "=SUM(C2:C2)"}

    out = photodoc.write_excel(data.rows, tmp_path / "ai.xlsx", data.title, data.merges)
    ws = openpyxl.load_workbook(out).active
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"
    assert ws["A2"].value.date().isoformat() == "2026-10-22"
    assert ws["C2"].value == 1250.5 and ws["C2"].number_format == "#,##0.00"
    assert list(ws.merged_cells.ranges)[0].coord == "A2:A3"
