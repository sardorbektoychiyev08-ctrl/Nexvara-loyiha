"""NEXVARA BI: loading, types, ETL steps, the automatic dashboard and its numbers."""

from __future__ import annotations

import datetime as dt
import json

import pytest

from nexvara.bi import dash
from nexvara.bi.data import (
    Table,
    apply_steps,
    eval_expr,
    load,
    parse_date,
    parse_number,
    typed,
)

SALES_CSV = """Sana;Filial;Mahsulot;Daromad;Xarajat
05.01.2026;Toshkent;Daftar;1 200 000;700 000
20.01.2026;Samarqand;Ruchka;800 000,50;500 000
03.02.2026;Toshkent;Ruchka;1 500 000;800 000
14.02.2026;Buxoro;Daftar;-;100 000
14.02.2026;Buxoro;Daftar;-;100 000
01.03.2026;Toshkent;Kitob;2 000 000;900 000
"""


@pytest.fixture
def sales(tmp_path):
    p = tmp_path / "savdo.csv"
    p.write_text(SALES_CSV, encoding="utf-8")
    return p


def test_numbers_and_dates():
    assert parse_number("1 234 567,5") == 1234567.5
    assert parse_number("1,234,567") == 1234567
    assert parse_number("1.234.567") == 1234567
    assert parse_number("12.5") == 12.5
    assert parse_number("250 000 so‘m") == 250000
    assert parse_number("15%") == 15
    assert parse_number("abc") is None and parse_number(True) is None
    assert parse_date("05.01.2026") == dt.date(2026, 1, 5)
    assert parse_date("2026-03") == dt.date(2026, 3, 1)
    assert parse_date("Mart 2026") == dt.date(2026, 3, 1)
    assert parse_date("salom") is None


def test_load_csv_types(sales):
    t = typed(load(sales))
    assert t.columns == ["Sana", "Filial", "Mahsulot", "Daromad", "Xarajat"]
    assert t.types == {"Sana": "date", "Filial": "text", "Mahsulot": "text", "Daromad": "number",
                       "Xarajat": "number"}  # fmt: skip
    assert t.rows[1][3] == 800000.5 and t.rows[3][3] is None


def test_load_xlsx_json_xml_folder(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Oylik hisobot"])  # a title above the table is skipped
    ws.append([])
    ws.append(["Oy", "Summa", None])
    ws.append([dt.datetime(2026, 1, 1), 100, None])
    ws.append([dt.datetime(2026, 2, 1), 150, None])
    wb.save(tmp_path / "a.xlsx")
    t = typed(load(tmp_path / "a.xlsx"))
    assert t.columns == ["Oy", "Summa"] and t.types["Oy"] == "date" and len(t.rows) == 2

    (tmp_path / "b.json").write_text(json.dumps({"data": [{"Oy": "2026-03", "Summa": 90}]}),
                                     encoding="utf-8")  # fmt: skip
    assert load(tmp_path / "b.json").rows == [["2026-03", 90]]
    (tmp_path / "c.xml").write_text(
        "<root><row id='1'><Oy>2026-04</Oy><Summa>70</Summa></row>"
        "<row id='2'><Oy>2026-05</Oy><Summa>60</Summa></row></root>", encoding="utf-8"
    )  # fmt: skip
    x = load(tmp_path / "c.xml")
    assert x.columns == ["id", "Oy", "Summa"] and len(x.rows) == 2
    folder = typed(load(tmp_path))
    assert folder.columns[0] == "Fayl" and len(folder.rows) == 5
    (tmp_path / "x.pdf").write_bytes(b"%PDF")
    with pytest.raises(ValueError, match="qo‘llanmaydi"):
        load(tmp_path / "x.pdf")


def test_expressions():
    row = {"Daromad": 100.0, "Xarajat": 40.0, "Soni": 3.0}
    assert eval_expr("[Daromad] - [Xarajat]", row) == 60
    assert eval_expr("([Daromad] - [Xarajat]) / [Soni] * 1,5", row) == 30
    assert eval_expr("-[Xarajat] + 10", row) == -30
    assert eval_expr("[Daromad] / 0", row) is None
    with pytest.raises(LookupError):
        eval_expr("[Yoq] + 1", row)
    with pytest.raises(ValueError):
        eval_expr("__import__('os')", row)


def test_etl_steps(sales, tmp_path):
    t = typed(load(sales))
    out = apply_steps(t, [
        {"op": "dedupe"},
        {"op": "fill_null", "col": "Daromad", "value": "__zero__"},
        {"op": "calc", "name": "Foyda", "expr": "[Daromad] - [Xarajat]"},
        {"op": "filter", "col": "Filial", "cmp": "!=", "value": "samarqand"},
        {"op": "rename", "col": "Filial", "name": "Shahar"},
        {"op": "drop_col", "col": "Mahsulot"},
    ])  # fmt: skip
    assert out.columns == ["Sana", "Shahar", "Daromad", "Xarajat", "Foyda"]
    assert len(out.rows) == 4 and out.rows[2][4] == -100000
    assert "1 ta takroriy" in out.notes[0]
    (tmp_path / "rahbar.csv").write_text(
        "Filial;Rahbar\nToshkent;Ali\nBuxoro;Vali\n", encoding="utf-8"
    )
    joined = apply_steps(t, [{"op": "join", "file": str(tmp_path / "rahbar.csv"), "on": "Filial"}])
    assert (
        joined.columns[-1] == "Rahbar"
        and joined.rows[0][-1] == "Ali"
        and joined.rows[1][-1] is None
    )
    assert (
        apply_steps(t, [{"op": "filter", "col": "Daromad", "cmp": ">", "value": "1 000 000"}]).rows[
            0
        ][1]
        == "Toshkent"
    )
    with pytest.raises(LookupError):
        apply_steps(t, [{"op": "drop_col", "col": "Yoq"}])


def test_auto_dashboard_profit_and_growth(sales):
    t = typed(load(sales))
    cfg = dash.auto_dashboard(t, str(sales))
    assert cfg["steps"] == [{"op": "calc", "name": "Sof foyda", "expr": "[Daromad] - [Xarajat]"}]
    assert cfg["date_col"] == "Sana" and "Filial" in cfg["slicers"]
    res = dash.compute(cfg)
    kpis = {w["title"]: w for w in res["widgets"] if w["type"] == "kpi"}
    assert kpis["Daromad (jami)"]["value"] == 5500000.5
    assert (
        kpis["Sof foyda (jami)"]["value"] == 5500000.5 - 2900000
    )  # rows without revenue have no profit
    assert kpis["Qatorlar soni"]["value"] == 6
    # March vs February revenue: 2 000 000 vs 1 500 000 -> +33.3 %
    assert kpis["Daromad (jami)"]["growth"] == pytest.approx(33.333, rel=1e-3)
    line = next(w for w in res["widgets"] if w["type"] == "line")
    assert line["labels"] == ["2026-01", "2026-02", "2026-03"]
    assert [s["name"] for s in line["series"]] == ["Daromad", "Xarajat", "Sof foyda"]
    bar = next(w for w in res["widgets"] if w["type"] == "bar")
    assert bar["labels"][0] == "Toshkent" and bar["values"][0] == 4700000
    assert res["date_range"] == ["2026-01-05", "2026-03-01"] and not res["problems"]


def test_filters_drilldown_and_pages(sales):
    cfg = dash.auto_dashboard(typed(load(sales)), str(sales))
    res = dash.compute(cfg, {"filters": {"Filial": ["Toshkent"]}})
    assert res["rows"] == 3 and res["total_rows"] == 6
    res = dash.compute(cfg, {"date_from": "2026-02-01", "date_to": "2026-02-28"})
    assert res["rows"] == 3
    res = dash.compute(cfg, {"page": 1})
    assert res["page"] == 1 and any(w["type"] == "scatter" for w in res["widgets"])
    table = next(w for w in res["widgets"] if w["type"] == "table")
    assert table["total_rows"] == 6 and table["rows"][0][0] == "2026-01-05"


def test_top_n_groups_rest():
    t = Table(["K", "V"], [[f"k{i}", float(i)] for i in range(10)], {"K": "text", "V": "number"})
    w = dash.compute_widget(t, {"type": "pie", "x": "K", "y": "V", "agg": "sum", "top": 4}, None)
    assert w["labels"] == ["k9", "k8", "k7", "Boshqalar"] and w["values"][-1] == sum(range(7))


def test_save_list_export(sales, tmp_path, monkeypatch):
    monkeypatch.setenv("NEXVARA_DATA", str(tmp_path / "data"))
    cfg = dash.auto_dashboard(typed(load(sales)), str(sales), name="Savdo paneli")
    path = dash.save(cfg)
    assert dash.saved()[0]["name"] == "Savdo paneli"
    assert dash.load_saved(str(path))["source"] == str(sales)
    with pytest.raises(ValueError):
        dash.load_saved(str(sales))
    csv_out = dash.export_rows(cfg, {"filters": {"Filial": ["Buxoro"]}}, "csv", tmp_path / "o.csv")
    assert csv_out.read_text(encoding="utf-8-sig").count("\n") == 3
    xlsx = dash.export_rows(cfg, {}, "xlsx", tmp_path / "o.xlsx")
    import openpyxl

    ws = openpyxl.load_workbook(xlsx).active
    assert ws["F1"].value == "Sof foyda" and ws.max_row == 7
    png = (
        "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8"
        "z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    assert dash.save_png(png, tmp_path / "a.png").stat().st_size > 0
    pdf = dash.pngs_to_pdf([png, png], tmp_path / "a.pdf", "Savdo")
    import pymupdf

    assert pymupdf.open(str(pdf)).page_count == 2
    with pytest.raises(ValueError):
        dash.save({**cfg, "name": ""})
