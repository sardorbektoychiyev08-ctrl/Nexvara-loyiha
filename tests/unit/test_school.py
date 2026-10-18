"""Teacher and office tools: answer sheets (photo → score), class journal, document
templates, weekly workflows, Wake-on-LAN info, remote-control guard rails, and their phone
routes. No real Word, OCR, mouse or keyboard is used: those parts are replaced by fakes."""

from __future__ import annotations

import random
import time
from pathlib import Path

import pytest
from test_remote_v1 import call, j, pair

from nexvara.core.task import UserError


def _photo(tmp_path, n, variants, answers, variant=1, angle=0.0, extra=None, name="photo.png"):
    """A filled answer sheet as a phone photo would show it (rotated, scaled)."""
    import fitz

    from nexvara import omr

    pdf = omr.make_sheet_pdf(tmp_path / "s.pdf", "Sinov", n, variants)
    doc = fitz.open(str(pdf))
    page = doc[0]
    lay = omr.layout(n, variants)
    for i, ch in enumerate(answers):
        if ch in "ABCD":
            x, y = lay["q"][i]["ABCD".index(ch)]
            page.draw_circle((x, y), omr.R * 0.8, color=(0.1, 0.1, 0.3), fill=(0.1, 0.1, 0.3))
    for i, k in extra or []:  # a second mark in the same row
        x, y = lay["q"][i][k]
        page.draw_circle((x, y), omr.R * 0.8, color=(0, 0, 0), fill=(0, 0, 0))
    if variants > 1:
        x, y = lay["variant"][variant - 1]
        page.draw_circle((x, y), omr.R * 0.8, color=(0, 0, 0), fill=(0, 0, 0))
    out = tmp_path / name
    page.get_pixmap(matrix=fitz.Matrix(1.7, 1.7).prerotate(angle)).save(str(out))
    doc.close()
    return out


@pytest.mark.parametrize("angle", [0, 8, -14])
def test_answer_sheet_is_read_from_a_rotated_photo(tmp_path, angle):
    from nexvara import omr

    rng = random.Random(angle)
    ans = "".join(rng.choice("ABCD") for _ in range(45))
    got = omr.read_sheet(_photo(tmp_path, 45, 3, ans, variant=3, angle=angle), 45, 3)
    assert got["answers"] == ans and got["variant"] == 3


def test_blank_and_double_marks_and_grading(tmp_path, monkeypatch):
    monkeypatch.setenv("NEXVARA_DATA", str(tmp_path / "data"))
    from openpyxl import load_workbook

    from nexvara import omr

    ans = "ABCD" * 5  # 20 questions
    marked = "-" + ans[1:]  # question 1 left empty
    photo = _photo(tmp_path, 20, 1, marked, extra=[(1, 3)])  # question 2: B and D
    got = omr.read_sheet(photo, 20, 1)
    assert got["answers"][0] == "-" and got["answers"][1] == "*" and got["answers"][2:] == ans[2:]
    s = omr.create_set("7-B fizika", 20, 1, {"1": ans})
    r = omr.grade(s["id"], photo, "Aliyev Vali")
    assert r["right"] == 18 and r["percent"] == 90.0 and r["mark"] == 5 and r["wrong"] == [1, 2]
    wb = load_workbook(r["excel"])
    rows = list(wb.active.values)
    assert rows[2][1] == "Aliyev Vali" and rows[2][3] == 18
    with pytest.raises(UserError):
        omr.set_keys(s["id"], {"1": "ABC"})  # wrong length
    with pytest.raises(UserError):
        omr.create_set("x", 61, 1)
    assert [omr.mark(p) for p in (90, 75, 60, 40)] == [5, 4, 3, 2]


def test_photo_without_corners_is_explained(tmp_path):
    import fitz

    from nexvara import omr

    doc = fitz.open()
    doc.new_page().insert_text((50, 50), "oddiy qog‘oz")
    p = tmp_path / "x.png"
    doc[0].get_pixmap().save(str(p))
    with pytest.raises(UserError, match="burchag"):
        omr.read_sheet(p, 10, 1)


def test_journal(tmp_path):
    from openpyxl import load_workbook

    from nexvara.journal import Journal

    jr = Journal(tmp_path / "app.db")
    cid = jr.add_class("7-B")
    with pytest.raises(UserError):
        jr.add_class("7-B")
    assert jr.add_students(cid, "1. Aliyev Vali\n2) Karimova Zuhra\n\nAliyev Vali") == 2
    v = jr.class_view(cid)
    vali, zuhra = v["students"]
    jr.add_grade(vali["id"], 5, "Fizika", "2026-10-15")
    jr.add_grade(vali["id"], 4, "Fizika", "2026-10-17")
    jr.add_grade(zuhra["id"], 3, "Algebra", "2026-10-17")
    with pytest.raises(UserError):
        jr.add_grade(vali["id"], 7, "Fizika")
    v = jr.class_view(cid, "Fizika")
    assert v["students"][0]["average"] == 4.5 and v["subjects"] == ["Algebra", "Fizika"]
    out = jr.export(cid, tmp_path)
    wb = load_workbook(out)
    assert set(wb.sheetnames) == {"Algebra", "Fizika"}
    rows = list(wb["Fizika"].values)
    assert rows[1][:4] == ("№", "F.I.Sh", "15.10", "17.10") and rows[2][1] == "Aliyev Vali"
    assert rows[2][-1] == 4.5
    jr.delete_class(cid)
    assert jr.classes() == []


def test_document_templates_render():
    from nexvara import doctemplates

    cat = doctemplates.catalog()
    assert {t["id"] for t in cat} >= {"ariza_tatil", "tushuntirish", "malumotnoma", "buyruq"}
    t = next(x for x in cat if x["id"] == "ariza_tatil")
    fields = {f["key"]: f["example"] for f in t["fields"]}
    blocks = doctemplates.render("ariza_tatil", fields)
    assert ("title", "ARIZA") in blocks and any("3 kunlik" in b for _k, b in blocks)
    with pytest.raises(UserError, match="To‘ldiring"):
        doctemplates.render("ariza_tatil", {**fields, "sabab": ""})
    items = doctemplates.render(
        "buyruq",
        {f["key"]: f["example"] for f in next(x for x in cat if x["id"] == "buyruq")["fields"]},
    )
    assert sum(1 for k, _ in items if k == "list") == 2


def test_weekly_workflow(tmp_path):
    from nexvara.core.task import TaskRunner
    from nexvara.workflows import Workflow, WorkflowEngine, WorkflowStore

    store = WorkflowStore(tmp_path / "wf.db")
    ran = []
    eng = WorkflowEngine(store, TaskRunner(), lambda text: ran.append(text) or {"reply": "ok"})
    with pytest.raises(ValueError):
        Workflow.from_dict(
            {
                "name": "x",
                "steps": [{"kind": "command", "text": "a"}],
                "trigger": {"type": "weekly", "days": [], "time": "08:00"},
            }
        ).validate()
    wf = store.save(
        Workflow.from_dict(
            {
                "name": "Dushanba hisobot",
                "steps": [{"kind": "command", "text": "hisobot"}],
                "trigger": {"type": "weekly", "days": [0], "time": "08:00"},
            }
        )
    )
    sunday = time.mktime((2026, 10, 18, 9, 0, 0, 0, 0, -1))
    monday_early = time.mktime((2026, 10, 19, 7, 0, 0, 0, 0, -1))
    monday = time.mktime((2026, 10, 19, 8, 5, 0, 0, 0, -1))
    assert eng.tick(sunday) == [] and eng.tick(monday_early) == []
    assert len(eng.tick(monday)) == 1 and eng.tick(monday) == []  # once a day
    assert store.get(wf.id).last_daily == "2026-10-19"


def test_wake_info_any_language():
    from nexvara.pcpower import wake_info

    en = (
        "Windows IP Configuration\n\nEthernet adapter Ethernet:\n\n   Physical Address. . . . : "
        "A4-BB-6D-12-34-56\n   IPv4 Address. . . . : 192.168.1.20(Preferred)\n   Subnet Mask . . : "
        "255.255.255.0

Wireless LAN adapter Wi-Fi:

   Physical Address. . : "
        "11-22-33-44-55-66
"
        "   IPv4 Address. . : 10.0.0.5(Preferred)\n   Subnet Mask . : 255.255.0.0\n"
    )
    ru = en.replace("Physical Address", "Физический адрес").replace("IPv4 Address", "IPv4-адрес")
    for text in (en, ru):
        assert wake_info("192.168.1.20", text) == {
            "mac": "A4:BB:6D:12:34:56",
            "broadcast": "192.168.1.255",
        }
        assert wake_info("10.0.0.5", text) == {
            "mac": "11:22:33:44:55:66",
            "broadcast": "10.0.255.255",
        }
    assert wake_info("172.16.0.1", en) is None


def test_control_rejects_bad_input_before_touching_anything():
    from nexvara import control

    for kwargs in (
        {"action": "click", "x": 1.5, "y": 0.2},
        {"action": "click"},
        {"action": "key", "key": "format c:"},
        {"action": "type", "text": ""},
        {"action": "shell"},
    ):
        with pytest.raises(ValueError):
            control.act(**kwargs)


# ---------------- phone routes ----------------
@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("NEXVARA_DATA", str(tmp_path / "data"))
    from nexvara.app.api import Api
    from nexvara.core.task import TaskRunner
    from nexvara.journal import Journal
    from nexvara.remote import DeviceStore

    done = {"control": [], "opened": [], "auto": False}

    def fake_doc(task, tid, blocks, out):
        out.write_text("\n".join(b for _k, b in blocks), encoding="utf-8")
        return {"path": str(out)}

    def fake_auto(on):
        done["auto"] = on
        return on

    api = Api(
        TaskRunner(),
        overrides={
            "device_store": DeviceStore(tmp_path / "app.db"),
            "remote_port": 0,
            "journal": Journal(tmp_path / "j.db"),
            "doctemplate": fake_doc,
            "ocr": lambda p, lang: "7-sinf fizika\nNyuton qonunlari",
            "control": lambda *a: done["control"].append(a),
            "office_pdf": lambda p: b"%PDF-fake",
            "autostart": fake_auto,
            "autostart_get": lambda: done["auto"],
        },
    )
    api.open_file = lambda p: done["opened"].append(p) or True
    assert api.devices_start()["ok"]
    srv = api._server()
    srv.host_url = f"https://127.0.0.1:{srv.port}"
    yield api, srv, pair(srv), done, tmp_path
    api.devices_stop()


def post(srv, path, data, token):
    return call(srv, "POST", path, data, token, {"Content-Type": "application/json"})


def _upload(srv, token, path: Path) -> str:
    import hashlib
    from urllib.parse import quote

    body = path.read_bytes()
    h = {"X-Filename": quote(path.name), "X-Content-SHA256": hashlib.sha256(body).hexdigest()}
    return j(call(srv, "POST", "/api/v1/files", token=token, raw=body, headers=h))["data"]["path"]


def test_routes_answer_sheets(env):
    api, srv, token, done, tmp = env
    ans = "ABCDA" * 4
    s = j(post(srv, "/api/v1/omr", {"title": "7-B", "questions": 20, "variants": 1,
                                    "keys": {"1": ans}}, token))["data"]  # fmt: skip
    assert s["questions"] == 20 and done["opened"] == [s["pdf"]]
    from urllib.parse import quote

    pdf = call(srv, "GET", f"/api/v1/files?path={quote(s['pdf'])}", token=token)
    assert pdf[0] == 200 and pdf[1][:4] == b"%PDF"  # printable sheet goes to the phone too
    photo = _upload(srv, token, _photo(tmp, 20, 1, ans))
    r = j(post(srv, f"/api/v1/omr/{s['id']}/grade", {"path": photo, "student": "Vali"}, token))[
        "data"
    ]
    assert r["right"] == 20 and r["mark"] == 5
    xl = call(srv, "GET", f"/api/v1/files?path={quote(r['excel'])}", token=token)
    assert xl[0] == 200
    lst = j(call(srv, "GET", "/api/v1/omr", token=token))["data"]
    assert lst["sets"][0]["results"][0]["student"] == "Vali" and lst["last_test"] is None
    outside = tmp / "x.png"
    outside.write_bytes(b"x")
    assert post(srv, f"/api/v1/omr/{s['id']}/grade", {"path": str(outside)}, token)[0] == 400
    assert post(srv, "/api/v1/omr", {"from_last_test": True}, token)[0] == 400
    # A test made from a textbook leaves its answer keys: the sheet uses them.
    api._runner.start("Test", lambda t: {"answers": ["ABCD", "DCBA"]}).wait(5)
    s2 = j(post(srv, "/api/v1/omr", {"from_last_test": True, "open_on_pc": False}, token))["data"]
    assert s2["questions"] == 4 and s2["variants"] == 2 and s2["keys"] == {"1": "ABCD", "2": "DCBA"}


def test_routes_journal_documents_ocr_preview(env):
    api, srv, token, done, tmp = env
    c = j(post(srv, "/api/v1/journal", {"name": "7-B"}, token))["data"]
    cid = c["id"]
    v = j(post(srv, f"/api/v1/journal/{cid}/students", {"names": "Ali\nVali"}, token))["data"]
    assert v["added"] == 2
    sid = v["students"][0]["id"]
    assert post(srv, "/api/v1/journal/grade", {"student_id": sid, "grade": 5, "subject": "Fizika"},
                token)[0] == 200  # fmt: skip
    assert post(srv, "/api/v1/journal/grade", {"student_id": sid, "grade": 9}, token)[0] == 400
    view = j(call(srv, "GET", f"/api/v1/journal/{cid}", token=token))["data"]
    assert view["students"][0]["average"] == 5.0
    x = j(post(srv, f"/api/v1/journal/{cid}/export", {"open_on_pc": False}, token))["data"]
    assert x["name"].endswith(".xlsx")
    from urllib.parse import quote

    assert call(srv, "GET", f"/api/v1/files?path={quote(x['path'])}", token=token)[0] == 200

    tpl = j(call(srv, "GET", "/api/v1/doctemplates", token=token))["data"]["templates"]
    t = next(x for x in tpl if x["id"] == "tushuntirish")
    fields = {f["key"]: f["example"] for f in t["fields"]}
    task = j(post(srv, "/api/v1/doctemplates/tushuntirish", {"fields": fields}, token))["data"][
        "task"
    ]
    assert api._runner.get(task["id"]).wait(5)
    assert post(srv, "/api/v1/doctemplates/tushuntirish", {"fields": {}}, token)[0] == 400
    assert post(srv, "/api/v1/doctemplates/nope", {"fields": {}}, token)[0] == 404

    img = tmp / "sahifa.jpg"
    img.write_bytes(b"\xff\xd8 page")
    up = _upload(srv, token, img)
    o = j(post(srv, "/api/v1/ocr", {"path": up}, token))["data"]
    assert "Nyuton" in o["text"] and api.get_context_file() == o["path"]
    pv = call(srv, "GET", f"/api/v1/files/preview?path={quote(o['path'])}", token=token)
    assert pv[0] == 200 and b"Nyuton" in pv[1] and pv[2]["Content-Type"].startswith("text/plain")
    xlsx = call(srv, "GET", f"/api/v1/files/preview?path={quote(x['path'])}", token=token)
    assert xlsx[0] == 200 and xlsx[1] == b"%PDF-fake"
    assert call(srv, "GET", f"/api/v1/files/preview?path={quote(str(tmp / 'nope.pdf'))}",
                token=token)[0] == 404  # fmt: skip


def test_routes_control_and_autostart(env):
    api, srv, token, done, _ = env
    r = post(srv, "/api/v1/control", {"action": "click", "x": 0.5, "y": 0.5}, token)
    assert r[0] == 400 and "o‘chiq" in j(r)["error"]["message"] and done["control"] == []
    api.devices_set_option("control", True)
    assert post(srv, "/api/v1/control", {"action": "key", "key": "enter"}, token)[0] == 200
    assert done["control"][0][0] == "key" and {"kind": "phone_control"} in api.ui_events()
    assert call(srv, "POST", "/api/v1/control", {"action": "click"})[0] == 401  # token needed
    assert j(post(srv, "/api/v1/autostart", {"on": True}, token))["data"] == {"on": True}
    assert j(call(srv, "GET", "/api/v1/autostart", token=token))["data"] == {"on": True}
    sess = j(call(srv, "GET", "/api/v1/session", token=token))["data"]
    assert "wake" in sess
