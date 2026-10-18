"""Class journal ("sinf jurnali"): classes, students, grades from the phone; Excel on the PC.

Stored in app.db (journal_* tables) next to memory; the Excel file is regenerated on demand
(one sheet per subject: students × dates, average, a column chart-ready layout).
"""

from __future__ import annotations

import re
import sqlite3
import threading
import time
from pathlib import Path
from typing import Any

from nexvara import config
from nexvara.core.task import UserError

_SCHEMA = """
CREATE TABLE IF NOT EXISTS journal_classes (
    id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE, created REAL NOT NULL);
CREATE TABLE IF NOT EXISTS journal_students (
    id INTEGER PRIMARY KEY, class_id INTEGER NOT NULL, name TEXT NOT NULL, pos INTEGER NOT NULL);
CREATE TABLE IF NOT EXISTS journal_grades (
    id INTEGER PRIMARY KEY, student_id INTEGER NOT NULL, subject TEXT NOT NULL, day TEXT NOT NULL,
    grade INTEGER NOT NULL, note TEXT NOT NULL DEFAULT '', created REAL NOT NULL);
"""


class Journal:
    def __init__(self, path: Path | None = None) -> None:
        self.path = path or (config.data_dir() / "app.db")
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()
        with self._db() as db:
            db.executescript(_SCHEMA)

    def _db(self) -> sqlite3.Connection:
        return sqlite3.connect(str(self.path), timeout=10)

    # ----- classes and students -----
    def classes(self) -> list[dict[str, Any]]:
        with self._db() as db:
            rows = db.execute(
                "SELECT c.id, c.name, COUNT(s.id) FROM journal_classes c "
                "LEFT JOIN journal_students s ON s.class_id = c.id GROUP BY c.id ORDER BY c.name"
            ).fetchall()
        return [{"id": r[0], "name": r[1], "students": r[2]} for r in rows]

    def add_class(self, name: str) -> int:
        name = re.sub(r"\s+", " ", (name or "").strip())[:30]
        if not name:
            raise UserError("Sinf nomini yozing, masalan «7-B»")
        with self._lock, self._db() as db:
            try:
                return int(
                    db.execute(
                        "INSERT INTO journal_classes (name, created) VALUES (?, ?)",
                        (name, time.time()),
                    ).lastrowid
                )
            except sqlite3.IntegrityError as e:
                raise UserError(f"«{name}» sinfi allaqachon bor") from e

    def delete_class(self, class_id: int) -> None:
        with self._lock, self._db() as db:
            ids = [
                r[0]
                for r in db.execute(
                    "SELECT id FROM journal_students WHERE class_id = ?", (class_id,)
                )
            ]
            db.executemany("DELETE FROM journal_grades WHERE student_id = ?", [(i,) for i in ids])
            db.execute("DELETE FROM journal_students WHERE class_id = ?", (class_id,))
            db.execute("DELETE FROM journal_classes WHERE id = ?", (class_id,))

    def add_students(self, class_id: int, names: str) -> int:
        """One name per line (a pasted list); numbering like "1. " is dropped."""
        clean = []
        for line in (names or "").splitlines():
            n = re.sub(r"^\s*\d+[.)]\s*", "", line).strip()
            n = re.sub(r"\s+", " ", n)[:60]
            if n:
                clean.append(n)
        if not clean:
            raise UserError("O‘quvchilar ismini yozing (har qatorda bittadan)")
        with self._lock, self._db() as db:
            if (
                db.execute("SELECT 1 FROM journal_classes WHERE id = ?", (class_id,)).fetchone()
                is None
            ):
                raise LookupError("Sinf topilmadi")
            pos = db.execute(
                "SELECT COALESCE(MAX(pos), 0) FROM journal_students WHERE class_id = ?", (class_id,)
            ).fetchone()[0]
            have = {
                r[0].casefold()
                for r in db.execute(
                    "SELECT name FROM journal_students WHERE class_id = ?", (class_id,)
                )
            }
            added = 0
            for n in clean:
                if n.casefold() in have:
                    continue
                have.add(n.casefold())
                pos += 1
                added += 1
                db.execute(
                    "INSERT INTO journal_students (class_id, name, pos) VALUES (?, ?, ?)",
                    (class_id, n, pos),
                )
        return added

    def remove_student(self, student_id: int) -> None:
        with self._lock, self._db() as db:
            db.execute("DELETE FROM journal_grades WHERE student_id = ?", (student_id,))
            db.execute("DELETE FROM journal_students WHERE id = ?", (student_id,))

    # ----- grades -----
    def add_grade(
        self, student_id: int, grade: int, subject: str, day: str = "", note: str = ""
    ) -> int:
        if int(grade) not in (1, 2, 3, 4, 5):
            raise UserError("Baho 1–5 bo‘lsin")
        subject = (subject or "").strip()[:40] or "Fan"
        day = day or time.strftime("%Y-%m-%d")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", day):
            raise UserError("Sana YYYY-MM-DD ko‘rinishida bo‘lsin")
        with self._lock, self._db() as db:
            if (
                db.execute("SELECT 1 FROM journal_students WHERE id = ?", (student_id,)).fetchone()
                is None
            ):
                raise LookupError("O‘quvchi topilmadi")
            return int(
                db.execute(
                    "INSERT INTO journal_grades (student_id, subject, day, grade, note, created) "
                    "VALUES (?, ?, ?, ?, ?, ?)",
                    (student_id, subject, day, int(grade), note.strip()[:100], time.time()),
                ).lastrowid
            )

    def delete_grade(self, grade_id: int) -> None:
        with self._lock, self._db() as db:
            db.execute("DELETE FROM journal_grades WHERE id = ?", (grade_id,))

    def class_view(self, class_id: int, subject: str = "") -> dict[str, Any]:
        with self._db() as db:
            c = db.execute(
                "SELECT id, name FROM journal_classes WHERE id = ?", (class_id,)
            ).fetchone()
            if c is None:
                raise LookupError("Sinf topilmadi")
            studs = db.execute(
                "SELECT id, name FROM journal_students WHERE class_id = ? ORDER BY pos", (class_id,)
            ).fetchall()
            subjects = [
                r[0]
                for r in db.execute(
                    "SELECT DISTINCT g.subject FROM journal_grades g JOIN journal_students s "
                    "ON s.id = g.student_id WHERE s.class_id = ? ORDER BY g.subject",
                    (class_id,),
                )
            ]
            sql = (
                "SELECT g.id, g.student_id, g.subject, g.day, g.grade FROM journal_grades g JOIN "
                "journal_students s ON s.id = g.student_id WHERE s.class_id = ?"
            )
            args: tuple = (class_id,)
            if subject:
                sql += " AND g.subject = ?"
                args += (subject,)
            grades = db.execute(sql + " ORDER BY g.day, g.id", args).fetchall()
        by: dict[int, list[dict[str, Any]]] = {}
        for g in grades:
            by.setdefault(g[1], []).append(
                {"id": g[0], "subject": g[2], "day": g[3], "grade": g[4]}
            )
        students = []
        for sid, name in studs:
            gs = by.get(sid, [])
            avg = round(sum(x["grade"] for x in gs) / len(gs), 2) if gs else None
            students.append(
                {"id": sid, "name": name, "grades": gs[-8:], "count": len(gs), "average": avg}
            )
        return {"id": c[0], "name": c[1], "subjects": subjects, "students": students}

    # ----- Excel -----
    def export(self, class_id: int, folder: Path | None = None) -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        view = self.class_view(class_id)
        wb = Workbook()
        wb.remove(wb.active)
        subjects = view["subjects"] or ["Baholar"]
        for subj in subjects:
            v = self.class_view(class_id, "" if subj == "Baholar" else subj)
            days = sorted({g["day"] for s in v["students"] for g in self._all(s["id"], subj)})
            ws = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", " ", subj)[:31])
            ws.append([f"{view['name']} · {subj}"])
            ws["A1"].font = Font(bold=True, size=14)
            ws.append(
                [
                    "№",
                    "F.I.Sh",
                    *[time.strftime("%d.%m", time.strptime(d, "%Y-%m-%d")) for d in days],
                    "O‘rtacha",
                ]
            )
            for c in ws[2]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="E3E4FF")
                c.alignment = Alignment(horizontal="center")
            for i, s in enumerate(v["students"], start=1):
                cells = {g["day"]: g["grade"] for g in self._all(s["id"], subj)}
                row = [i, s["name"], *[cells.get(d, "") for d in days]]
                gs = [x for x in self._all(s["id"], subj)]
                row.append(round(sum(x["grade"] for x in gs) / len(gs), 2) if gs else "")
                ws.append(row)
            ws.column_dimensions["A"].width = 5
            ws.column_dimensions["B"].width = 30
        out_dir = folder or (config.data_dir() / "Jurnal")
        out_dir.mkdir(parents=True, exist_ok=True)
        safe = re.sub(r"[^\w -]", "", view["name"])
        out = out_dir / f"Jurnal {safe}.xlsx"
        wb.save(out)
        return out

    def _all(self, student_id: int, subject: str) -> list[dict[str, Any]]:
        with self._db() as db:
            sql = "SELECT day, grade FROM journal_grades WHERE student_id = ?"
            args: tuple = (student_id,)
            if subject and subject != "Baholar":
                sql += " AND subject = ?"
                args += (subject,)
            return [
                {"day": r[0], "grade": r[1]} for r in db.execute(sql + " ORDER BY day, id", args)
            ]
