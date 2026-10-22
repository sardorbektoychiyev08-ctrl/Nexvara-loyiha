"""A photo of a paper table or document → Excel or Word («shu rasmni Excel qil»).

The text comes from the free offline OCR (or, after the user's «Ha», from a vision AI, which
writes table cells with ' | '). Rows are found from ' | ', tabs or wide gaps; a table becomes
an .xlsx with a bordered, bold header (openpyxl, no Excel window needed) or a Word table;
plain text becomes Word paragraphs. The result is saved next to the photo as
«<rasm>_NEXVARA.xlsx/.docx» — the photo is never changed.
"""

from __future__ import annotations

import datetime as dt
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from nexvara.core.task import UserError

_SPLIT = re.compile(r"\s*\|\s*|\t+| {3,}")
_NUM = re.compile(r"^[+-]?\d{1,3}(?:[  ]?\d{3})*(?:[.,]\d+)?$|^[+-]?\d+(?:[.,]\d+)?$")
MAX_ROWS = 5000
_DATE = re.compile(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})$|^(\d{1,2})[./-](\d{1,2})[./-](\d{4})$")
_CELL = re.compile(r"^[A-Z]{1,3}[1-9]\d*$")


@dataclass
class TableData:
    rows: list[list[Any]]
    title: str = "Jadval"
    merges: list[str] = field(default_factory=list)
    formulas: dict[str, str] = field(default_factory=dict)


def _date(cell: str) -> dt.date | None:
    m = _DATE.fullmatch(cell.strip())
    if not m:
        return None
    try:
        if m.group(1):
            return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        return dt.date(int(m.group(6)), int(m.group(5)), int(m.group(4)))
    except ValueError:
        return None


def value(cell: Any, *, header: bool = False) -> Any:
    if not isinstance(cell, str) or header:
        return cell
    return _date(cell) or number(cell)


def _clean_rows(rows: Any) -> list[list[Any]]:
    if not isinstance(rows, list):
        return []
    clean = [list(r) for r in rows if isinstance(r, list) and any(str(c).strip() for c in r)]
    if len(clean) < 2:
        return []
    width = max(len(r) for r in clean)
    clean = [(r + [""] * width)[:width] for r in clean[:MAX_ROWS]]
    keep = [c for c in range(width) if any(str(r[c]).strip() for r in clean)]
    clean = [[r[c] for c in keep] for r in clean]
    return [[value(c, header=i == 0) for c in r] for i, r in enumerate(clean)]


def parse_ai_table(text: str) -> TableData:
    """Validated vision output. A pipe table is accepted for older/fallback providers."""
    raw = (text or "").strip()
    data: Any = None
    a, b = raw.find("{"), raw.rfind("}")
    if 0 <= a < b:
        try:
            data = json.loads(raw[a : b + 1])
        except ValueError:
            data = None
    if isinstance(data, dict):
        rows = _clean_rows(data.get("rows"))
        title = str(data.get("title") or "Jadval").strip()[:60] or "Jadval"
        merges = []
        for address in data.get("merges") or []:
            address = str(address).upper().replace("$", "")
            if re.fullmatch(r"[A-Z]{1,3}[1-9]\d*:[A-Z]{1,3}[1-9]\d*", address):
                merges.append(address)
        formulas = (
            {
                str(k).upper().replace("$", ""): str(v)
                for k, v in (data.get("formulas") or {}).items()
                if _CELL.fullmatch(str(k).upper().replace("$", "")) and str(v).startswith("=")
            }
            if isinstance(data.get("formulas"), dict)
            else {}
        )
        if rows:
            return TableData(rows, title, merges, formulas)
    rows = table_rows(raw)
    if rows:
        return TableData(_clean_rows(rows))
    raise UserError("AI rasmda jadval topmadi. Rasmni yaqinroq va tekisroq suratga oling.")


def split_row(line: str) -> list[str]:
    line = line.strip().strip("|").strip()
    if not line:
        return []
    return [c.strip() for c in _SPLIT.split(line)]


def table_rows(text: str) -> list[list[str]] | None:
    """The table in OCR text, or None when it is ordinary text. Lines that are only
    separators (---|---) are skipped; short rows are padded to the widest common row."""
    rows = []
    for line in (text or "").splitlines():
        if re.fullmatch(r"[\s|:+\-=_]*", line):
            continue
        cells = split_row(line)
        if cells:
            rows.append(cells)
    multi = [r for r in rows if len(r) >= 2]
    if len(multi) < 2 or len(multi) < 0.6 * len(rows):
        return None
    width = max(len(r) for r in multi)
    # a one-cell first line is the title above the table; later one-cell lines are kept
    out = [r + [""] * (width - len(r)) for i, r in enumerate(rows) if len(r) >= 2 or i > 0]
    return out[:MAX_ROWS]


def number(cell: str) -> Any:
    """«1 250 000» / «12,5» → a real number for Excel; anything else stays text
    (leading zeros like «007» and phone-like numbers stay text)."""
    s = cell.strip()
    if not s or not _NUM.match(s) or (len(s) > 1 and s[0] == "0" and s[1].isdigit()):
        return cell
    s = s.replace(" ", "").replace(" ", "")
    if len(re.sub(r"\D", "", s)) > 15:
        return cell
    try:
        return int(s) if re.fullmatch(r"[+-]?\d+", s) else float(s.replace(",", "."))
    except ValueError:
        return cell


def write_excel(
    rows: list[list[Any]],
    path: Path,
    title: str = "Jadval",
    merges: list[str] | None = None,
) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from nexvara.skills.excel.bridge import BORDER_COLOR, HEADER_FILL, HEADER_FONT

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "", title)[:31] or "Jadval"
    thin = Side(style="thin", color=BORDER_COLOR.lstrip("#"))
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r, row in enumerate(rows, start=1):
        for c, cell in enumerate(row, start=1):
            v = value(cell, header=r == 1)
            x = ws.cell(row=r, column=c, value=v)
            x.border = border
            x.alignment = Alignment(vertical="center", wrap_text=r == 1,
                                    horizontal="center" if r == 1 or not isinstance(v, str)
                                    else "left")  # fmt: skip
            if r == 1:
                x.font = Font(bold=True, color=HEADER_FONT.lstrip("#"))
                x.fill = PatternFill("solid", fgColor=HEADER_FILL.lstrip("#"))
            elif isinstance(v, (dt.date, dt.datetime)):
                x.number_format = "dd.mm.yyyy"
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                x.number_format = "#,##0.00" if isinstance(v, float) else "#,##0"
    for c in range(1, len(rows[0]) + 1):
        longest = max(len(str(row[c - 1])) for row in rows[:300])
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = min(
            max(8, longest + 3), 60
        )
    ws.freeze_panes = "A2"
    for address in merges or []:
        ws.merge_cells(address)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _kind(v: Any) -> str:
    if isinstance(v, (dt.date, dt.datetime)):
        return "date"
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return ""
    return "integer" if isinstance(v, int) else "decimal"


def write_excel_com(data: TableData, path: Path, keep_open: bool = True) -> Path:
    """Create the visible workbook with the project's locale-safe Excel helpers."""
    from nexvara.skills.excel import ops
    from nexvara.skills.excel.bridge import ExcelSession, finish_table, freeze_header

    with ExcelSession(visible=True, keep_open=keep_open) as excel:
        wb = excel.new_workbook()
        ws = wb.Worksheets(1)
        ws.Name = re.sub(r"[\[\]:*?/\\]", "", data.title)[:31] or "Jadval"
        rng = ops.write_rows(ws, data.rows)
        excel.app.DisplayAlerts = False  # «only the upper-left value is kept» would block
        try:
            for address in data.merges:
                ws.Range(address).Merge()
        finally:
            excel.app.DisplayAlerts = True
        for address, formula in data.formulas.items():
            ws.Range(address).Formula = formula
        # Formats go by each cell's type: a «Jami» row or a note must not turn the dates of
        # its column into serial numbers (Office check 2026-10-22: 46317 instead of 22.10.2026).
        kinds = [[_kind(v) for v in r] for r in data.rows[1:]]
        numeric = [c + 1 for c in range(len(data.rows[0]))
                   if any(k[c] in ("integer", "decimal") for k in kinds)]  # fmt: skip
        finish_table(ws, rng, numeric, 1, 1)
        for c in range(len(data.rows[0])):
            col = [k[c] for k in kinds]
            whole = "decimal" not in col
            r = 0
            while r < len(col):  # one COM call per run of same-kind cells
                kind = col[r]
                end_r = r
                while end_r + 1 < len(col) and col[end_r + 1] == kind:
                    end_r += 1
                if kind:
                    if kind == "integer" and not whole:
                        kind = "decimal"
                    ops.set_number_format(ws.Range(ws.Cells(r + 2, c + 1),
                                                   ws.Cells(end_r + 2, c + 1)), kind)  # fmt: skip
                r = end_r + 1
        freeze_header(excel.app, 1)
        excel.show()
        saved = excel.save_as(wb, path)
        del rng, ws, wb
    return saved


def write_word(text: str, rows: list[list[str]] | None, path: Path, title: str,
               visible: bool = True, keep_open: bool = True) -> Path:  # fmt: skip
    from nexvara.skills.word import ops

    with ops.WordSession(visible=visible, keep_open=keep_open) as w:
        doc = w.new_document()
        ops.add_paragraph(doc, title, ops.STYLE_HEADING[1], align="center")
        if rows:
            clean = [[c.replace("\t", " ") for c in r] for r in rows]
            ops.add_table(doc, clean)
        else:
            for para in re.split(r"\n\s*\n", text.strip()):
                line = " ".join(para.split())
                if line:
                    ops.add_paragraph(doc, line)
        w.show()
        saved = w.save_as(doc, path)
        del doc
    return saved


def output_for(image: str | Path, ext: str) -> Path:
    from nexvara.agent.files import output_path

    src = Path(image)
    if src.suffix.lower() == ".txt":  # the OCR text file next to the photo
        src = src.with_suffix(".jpg")
    return output_path(src, ext)


def convert(
    text: str | TableData,
    image: str | Path,
    to: str,
    task: Any = None,
    visible: bool = True,
    keep_open: bool = True,
    use_com: bool = False,
) -> dict[str, Any]:
    if isinstance(text, TableData):
        data = text
        body = "\n".join(" | ".join(str(c) for c in r) for r in data.rows)
    else:
        body = text
        data = None
    if not (body or "").strip():
        raise UserError("Rasmda matn topilmadi: yaqinroq, yorug‘ joyda, tekis suratga oling.")
    rows = data.rows if data else table_rows(body)
    title = Path(image).stem.replace("_", " ")[:60]
    if to == "excel":
        if rows is None:
            raise UserError(
                "Rasmda jadval topilmadi (ustunlar ajralmadi). Jadvalli rasm uchun "
                "«AI bilan Excel qil» deb yozing yoki Word’ga aylantiring."
            )
        data = data or TableData(_clean_rows(rows), title)
        out_path = output_for(image, ".xlsx")
        out = (
            write_excel_com(data, out_path, keep_open=keep_open)
            if use_com
            else write_excel(data.rows, out_path, data.title or title, data.merges)
        )
        return {"path": str(out), "rows": len(rows) - 1, "cols": len(rows[0]), "kind": "excel"}
    if to == "word":
        out = write_word(body, rows, output_for(image, ".docx"), title, visible, keep_open)
        return {"path": str(out), "rows": len(rows) - 1 if rows else 0, "kind": "word"}
    raise UserError("Faqat Excel yoki Word")
