"""Reading a photo in chat (mixed into app.api.Api): "shu rasmdagi matnni o‘qi".

Free and offline by default (Windows OCR, enhanced photo, best of Latin/Cyrillic). With
"AI bilan" — handwriting, tables, messy photos — the picture goes to a vision model after the
user's "Ha" (EXIF/GPS removed first, like drawings); the text is saved next to the photo.
"""

from __future__ import annotations

import datetime as dt
import json
import re
import uuid
from pathlib import Path
from typing import Any

_READ = re.compile(
    r"\b(matn\w*|yozuv\w*|o[‘’']?qi\w*|ko[‘’']?chir\w*|ocr|tekst\w*|qo[‘’']?lyozma\w*)\b",
    re.IGNORECASE,
)
PROMPT = (
    "Transcribe ALL text in the image exactly as written (printed or handwritten), in its own "
    "language and script (Uzbek Latin/Cyrillic, Russian, English). Keep paragraphs and line "
    "breaks, render tables as lines with ' | ' between cells. Output only the text."
)
TABLE_PROMPT = """Rasmda ko‘ringan jadvalni aniq ko‘chir. Faqat JSON qaytar:
{"title":"qisqa nom","rows":[["Sarlavha 1","Sarlavha 2"],["qiymat",12.5]],
 "merges":["A2:B2"],"formulas":{"C5":"=SUM(C2:C4)"}}
Qoidalar: rows[0] ustun sarlavhalari; rasmda yo‘q ma’lumotni to‘qima; keraksiz bo‘sh qator va
butunlay bo‘sh ustunni chiqarma; sonni JSON soni, sanani YYYY-MM-DD yoz; formulani faqat rasmda
formula yoki jami hisoblash aniq ko‘rinsa ber; birlashgan kataklarni merges’da A1 manzilida ber.
Bugungi sana: {today}. Foydalanuvchi talabi: {request}"""


def _table_prompt(request: str) -> str:
    return TABLE_PROMPT.replace("{today}", dt.date.today().isoformat()).replace(
        "{request}", request.strip() or "jadvalning hammasi"
    )


_TO_OFFICE = re.compile(
    r"\b(excel|eksel|xlsx|jadval\w*|word|vord|docx|hujjat\w*)\b.*\b(qil|aylantir|otkaz|ko.?chir|"
    r"o.?tkaz|yarat|tuz|sol)\w*|\b(qil|aylantir|o.?tkaz)\w*\b.*\b(excel|word)\w*",
    re.IGNORECASE,
)


def office_target(text: str) -> str | None:
    """«shu rasmni Excel qil» -> "excel", «… Word hujjatga aylantir» -> "word"."""
    from nexvara.knowledge.textnorm import fold

    f = fold(text)
    if not _TO_OFFICE.search(f):
        return None
    if re.search(r"\b(word|vord|docx|hujjat\w*)\b", f):
        return "word"
    return "excel"


class ReadingApi:
    _overrides: dict
    _pending: dict

    def _chat_read(self, text: str) -> dict[str, Any] | None:
        from nexvara import drawing
        from nexvara.app.api import _wants_ai

        cur = self._context_file
        from nexvara.agent.parse import find_file

        to = office_target(text or "") if not find_file(text or "") else None
        wants_ai = _wants_ai(text) and not bool(getattr(self._local, "ai_fallback", False))
        if to and cur and (drawing.is_image(cur) or (cur.lower().endswith(".txt")
                                                      and Path(cur).is_file())):  # fmt: skip
            return self._chat_photo_office(cur, to, wants_ai, text)
        if (
            not drawing.is_image(cur)
            or not _READ.search(text or "")
            or drawing.is_drawing_request(text)
        ):
            return None
        if wants_ai:
            if not self._ai_ready():
                return {
                    "reply": "AI bilan o‘qish uchun Sozlamalar → AI’da OpenAI, Gemini yoki Claude "
                    "kalitini kiriting. AI’siz o‘qish uchun «AI bilan» so‘zini yozmang."
                }
            if self._vision_provider() is None:
                return {
                    "reply": "Rasmni o‘qiydigan AI ulanmagan: OpenAI, Gemini yoki Claude kerak."
                }
            from nexvara.ai.router import AiUnavailable

            msgs = [
                {
                    "role": "user",
                    "content": "Read the image.",
                    "images": [{"mime": "image/jpeg", "data": ""}],
                }
            ]
            try:
                est = self._router().estimate(
                    "balanced", PROMPT, msgs, 1500, self._vision_provider()
                )
            except AiUnavailable as e:
                return {"reply": str(e)}
            import uuid

            token = uuid.uuid4().hex[:12]
            self._pending[token] = ("read", {"image": cur})
            return {
                "reply": f"Rasmdagi matnni AI o‘qiydi (qo‘lyozma ham): {est.describe()}."
                "\n\n📷 Rasm AI’ga yuboriladi (joylashuv ma’lumotlari o‘chiriladi). "
                "Unda shaxsiy ma’lumot bo‘lmasin."
                "\n\nDavom etaymi?",
                "confirm": token,
            }
        reader = self._overrides.get("ocr")
        if reader is None:
            from nexvara.knowledge.ocr import ocr_image as reader
        lang = "ru" if re.search(r"\bkirill\w*|\brus\w*", text, re.I) else ""
        try:
            out = reader(Path(cur), lang)
        except Exception as e:
            return {"reply": f"Rasmni o‘qib bo‘lmadi: {e}"}
        return self._read_done(cur, out, "kompyuterda, internetsiz")

    def _chat_photo_office(self, cur: str, to: str, ai: bool, text: str) -> dict[str, Any]:
        """Photo (or its OCR text) → Excel/Word. With «AI bilan» the photo is read by a
        vision AI after «Ha» (tables come out much better)."""
        from nexvara import drawing

        if ai and drawing.is_image(cur):
            return self._ask_read_ai(cur, to, text)
        if drawing.is_image(cur):
            reader = self._overrides.get("ocr")
            if reader is None:
                from nexvara.knowledge.ocr import ocr_image as reader
            lang = "ru" if re.search(r"\bkirill\w*|\brus\w*", text, re.I) else ""
            try:
                body = reader(Path(cur), lang)
            except Exception as e:
                return {"reply": f"Rasmni o‘qib bo‘lmadi: {e}"}
            Path(cur).with_suffix(".txt").write_text(body, encoding="utf-8")
            how = "kompyuterda, internetsiz"
        else:
            body = Path(cur).read_text(encoding="utf-8", errors="replace")
            how = "oldin o‘qilgan matndan"
        return self._photo_to_office(cur, body, to, how)

    def _photo_to_office(self, image: str, body: Any, to: str, how: str) -> dict[str, Any]:
        from nexvara import photodoc
        from nexvara.core.task import UserError

        if to == "excel":
            use_com = self._history is not None

            def work(task: Any) -> dict[str, Any]:
                with task.step("Excel jadval") as step:
                    made = photodoc.convert(
                        body, image, "excel", keep_open=self._keep_open(), use_com=use_com
                    )
                    step.detail = Path(made["path"]).name
                    return made

            if use_com:
                snap = self._runner.start("Rasm → Excel", work).snapshot()
                return {
                    "reply": f"📊 Rasmdagi jadvalni Excel’ga o‘tkazyapman ({how}) — Excel "
                    "kompyuter ekranida ochiladi. Jarayon o‘ng panelda.",
                    "task": snap,
                }
            try:
                res = photodoc.convert(body, image, "excel")
            except UserError as e:
                return {"reply": str(e), "error": True}
            self._ctx.result = res["path"]
            self._context_file = res["path"]
            return {"reply": f"📊 Rasmdagi jadval Excel’ga o‘tkazildi ({how}): "
                    f"{res['rows']} qator, "
                    f"{res['cols']} ustun, ramka va sarlavha bilan.\nFayl: {res['path']}\n"
                    "Raqamlarni rasm bilan solishtirib chiqing: OCR ba’zan adashadi.",
                    "files": [res["path"]]}  # fmt: skip
        fn = self._overrides.get("photo_word")
        keep = self._keep_open()

        def work(task: Any) -> dict[str, Any]:
            with task.step("Word hujjat") as s:
                if fn is not None:
                    res = fn(body, image)
                else:
                    res = photodoc.convert(body, image, "word", keep_open=keep)
                s.detail = Path(res["path"]).name
            return res

        snap = self._runner.start("Rasm → Word", work).snapshot()
        return {"reply": f"📄 Rasmdagi matnni ({how}) Word hujjatga yozyapman — jarayon o‘ng "
                "panelda.", "task": snap}  # fmt: skip

    def _ask_read_ai(self, cur: str, to: str = "", request: str = "") -> dict[str, Any]:
        if not self._ai_ready():
            return {
                "reply": "AI bilan o‘qish uchun Sozlamalar → AI’da OpenAI, Gemini yoki Claude "
                "kalitini kiriting. AI’siz o‘qish uchun «AI bilan» so‘zini yozmang."
            }
        if self._vision_provider() is None:
            return {"reply": "Rasmni o‘qiydigan AI ulanmagan: OpenAI, Gemini yoki Claude kerak."}
        from nexvara.ai.router import AiUnavailable

        msgs = [{"role": "user", "content": "Read the image.",
                 "images": [{"mime": "image/jpeg", "data": ""}]}]  # fmt: skip
        prompt = _table_prompt(request) if to == "excel" else PROMPT
        try:
            est = self._router().estimate("powerful", prompt, msgs, 1500, self._vision_provider())
        except AiUnavailable as e:
            return {"reply": str(e)}
        import uuid

        token = uuid.uuid4().hex[:12]
        self._pending[token] = ("read", {"image": cur, "to": to, "request": request})
        what = {"excel": " va Excel jadvalga o‘tkazadi", "word": " va Word hujjatga yozadi"}
        return {"reply": f"Rasmdagi matnni AI o‘qiydi (qo‘lyozma ham){what.get(to, '')}: "
                f"{est.describe()}.\n\n📷 Rasm AI’ga yuboriladi (joylashuv ma’lumotlari "
                "o‘chiriladi). Unda shaxsiy ma’lumot bo‘lmasin.\n\nDavom etaymi?",
                "confirm": token}  # fmt: skip

    def _read_done(self, image: str, out: str, how: str) -> dict[str, Any]:
        txt = Path(image).with_suffix(".txt")
        txt.write_text(out, encoding="utf-8")
        self._context_file = str(txt)  # "shu matndan test tuz" continues from here
        if not out.strip():
            return {
                "reply": "Rasmda matn topilmadi. Yaqinroq, yorug‘ joyda, tekis suratga oling yoki "
                "«AI bilan o‘qi» deb yozing (qo‘lyozma uchun)."
            }
        body = out if len(out) <= 3500 else out[:3500] + "…"
        return {
            "reply": f"📄 Rasmdagi matn ({how}):\n\n{body}\n\nFayl: {txt.name}. Endi «shu matndan "
            "test tuz» yoki «AI bilan qisqa konspekt qil» deyishingiz mumkin.",
            "text_file": str(txt),
        }

    def _run_read(self, payload: dict[str, Any]) -> dict[str, Any]:
        from nexvara import drawing
        from nexvara.core.task import UserError

        img = drawing.clean_image(Path(payload["image"]), max_side=2000)
        msgs = [
            {
                "role": "user",
                "content": "Read the image.",
                "images": [{"mime": img[0], "data": img[1]}],
            }
        ]
        prompt = (
            _table_prompt(payload.get("request", "")) if payload.get("to") == "excel" else PROMPT
        )
        try:
            reply = self._router().complete(
                "ocr",
                prompt,
                msgs,
                max_tokens=4000,
                expected_out=1500,
                approved=True,
                tier="powerful",
            )
        except UserError as e:
            return {"reply": str(e)}
        how = f"AI, ≈${reply.cost:.3f}"
        if payload.get("to"):
            if payload["to"] == "excel":
                from nexvara import photodoc

                body = photodoc.parse_ai_table(reply.text)
                saved_text = "\n".join(" | ".join(str(c) for c in row) for row in body.rows)
            else:
                body = reply.text.strip()
                saved_text = body
            Path(payload["image"]).with_suffix(".txt").write_text(saved_text, encoding="utf-8")
            return self._photo_to_office(payload["image"], body, payload["to"], how)
        return self._read_done(payload["image"], reply.text.strip(), how)

    def _ask_compare_table(self, request: str = "") -> dict[str, Any]:
        image = self._ctx.image
        excel = self._ctx.result if self._ctx.result.lower().endswith(".xlsx") else self._ctx.file
        if not image or not Path(image).is_file():
            return {"reply": "Solishtirish uchun asl rasm topilmadi.", "error": True}
        if not excel or not excel.lower().endswith(".xlsx") or not Path(excel).is_file():
            return {"reply": "Solishtirish uchun yaratilgan Excel fayl topilmadi.", "error": True}
        if not self._ai_ready() or self._vision_provider() is None:
            return {"reply": "Rasm va Excel’ni solishtirish uchun rasmni ko‘radigan AI kerak."}
        rows = _workbook_rows(Path(excel))
        system = _compare_prompt(rows, request)
        msgs = [{"role": "user", "content": "Rasmni Excel jadvali bilan tekshir.",
                 "images": [{"mime": "image/jpeg", "data": ""}]}]  # fmt: skip
        from nexvara.ai.router import AiUnavailable

        try:
            est = self._router().estimate("powerful", system, msgs, 2000, self._vision_provider())
        except AiUnavailable as e:
            return {"reply": str(e)}
        token = uuid.uuid4().hex[:12]
        self._pending[token] = (
            "compare_table",
            {"image": image, "excel": excel, "request": request, "rows": rows},
        )
        return {
            "reply": f"Rasmni yaratilgan Excel bilan katakma-katak solishtiraman: "
            f"{est.describe()}.\n\n📷 Asl rasm AI’ga yuboriladi. Davom etaymi?",
            "confirm": token,
        }

    def _run_compare_table(self, payload: dict[str, Any]) -> dict[str, Any]:
        from nexvara import drawing, photodoc
        from nexvara.core.task import UserError

        img = drawing.clean_image(Path(payload["image"]), max_side=2000)
        system = _compare_prompt(payload["rows"], payload.get("request", ""))
        msgs = [{"role": "user", "content": "Rasmni Excel jadvali bilan tekshir.",
                 "images": [{"mime": img[0], "data": img[1]}]}]  # fmt: skip
        try:
            reply = self._router().complete(
                "vision",
                system,
                msgs,
                max_tokens=5000,
                expected_out=2000,
                approved=True,
                tier="powerful",
            )
            data = _json_object(reply.text)
        except (UserError, ValueError) as e:
            return {"reply": str(e)}
        diffs = [str(x).strip() for x in data.get("differences") or [] if str(x).strip()]
        if not diffs:
            return {"reply": "Rasm va Excel jadvali mos: xato topilmadi."}
        corrected = photodoc.parse_ai_table(
            json.dumps(data.get("corrected") or {}, ensure_ascii=False)
        )
        token = uuid.uuid4().hex[:12]
        self._pending[token] = (
            "photo_fix",
            {"image": payload["image"], "table": corrected},
        )
        return {
            "reply": "Topilgan farqlar:\n"
            + "\n".join(f"• {x}" for x in diffs[:20])
            + "\n\nTuzatilganini yangi Excel fayl qilib, ekranda ochaymi?",
            "confirm": token,
        }

    def _run_photo_fix(self, payload: dict[str, Any]) -> dict[str, Any]:
        from nexvara import photodoc

        data = payload["table"]
        use_com = self._history is not None
        res = photodoc.convert(
            data, payload["image"], "excel", keep_open=self._keep_open(), use_com=use_com
        )
        self._ctx.result = res["path"]
        self._context_file = res["path"]
        return {
            "reply": f"Tuzatilgan yangi Excel tayyor va ekranda ochildi: {res['path']}",
            "files": [res["path"]],
        }


def _workbook_rows(path: Path) -> list[list[Any]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb.active
        return [list(r) for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200),
                                               values_only=True)]  # fmt: skip
    finally:
        wb.close()


def _compare_prompt(rows: list[list[Any]], request: str) -> str:
    return (
        "Rasmdagi jadvalni quyidagi Excel qiymatlari bilan aniq solishtir. Faqat JSON qaytar: "
        '{"differences":["B3: rasmda 15, Excelda 51"],"corrected":'
        '{"title":"Jadval","rows":[["Sarlavha"],["qiymat"]],"merges":[],"formulas":{}}}. '
        "Xato bo‘lmasa differences bo‘sh bo‘lsin. Rasmda ko‘rinmagan narsani to‘qima. "
        f"Talab: {request or 'hamma katakni tekshir'}. Excel: "
        + json.dumps(rows, ensure_ascii=False, default=str)
    )


def _json_object(text: str) -> dict[str, Any]:
    a, b = text.find("{"), text.rfind("}")
    data = json.loads(text[a : b + 1]) if 0 <= a < b else None
    if not isinstance(data, dict):
        raise ValueError("AI solishtirish javobini tushunib bo‘lmadi.")
    return data
