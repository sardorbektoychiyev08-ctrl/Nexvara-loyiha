"""Dashboards: an automatic first design for any table, the numbers behind every widget
(with slicers, a date range and drill-down filters), saving and exporting.

A dashboard config is plain JSON:
{"name", "source", "sheet", "steps": [...ETL...], "slicers": [cols], "date_col",
 "pages": [{"title", "widgets": [{"id", "type", "title", "x", "y", "agg", "period", "top"}]}]}
"""

from __future__ import annotations

import datetime as dt
import json
import re
import time
from collections import defaultdict
from pathlib import Path
from typing import Any

from nexvara import config as app_config
from nexvara.bi.data import Table, apply_steps, load, typed
from nexvara.knowledge.textnorm import fold

WIDGET_TYPES = ("kpi", "bar", "line", "pie", "scatter", "table")
AGGS = ("sum", "avg", "count", "min", "max")
AGG_LABELS = {
    "sum": "jami",
    "avg": "o‘rtacha",
    "count": "soni",
    "min": "eng kichik",
    "max": "eng katta",
}
PERIODS = ("day", "month", "quarter", "year")

_REVENUE = ("daromad", "tushum", "savdo", "sotuv", "summa", "vyruchk", "revenue", "sales", "income",
            "выручка", "доход", "продаж")  # fmt: skip
_EXPENSE = ("xarajat", "harajat", "chiqim", "tannarx", "rasxod", "expense", "cost", "расход",
            "затрат")  # fmt: skip
_ID_LIKE = re.compile(r"(^|\b)(id|№|no|nomer|raqam\w*|kod\w*|tr|t/r)(\b|$)")


# ---------------- automatic design ----------------


def _find(cols: list[str], keys: tuple[str, ...]) -> str | None:
    for c in cols:
        f = fold(c)
        if any(k in f or k in c.lower() for k in keys):
            return c
    return None


def profile(t: Table) -> dict[str, Any]:
    nums = [c for c in t.columns if t.types.get(c) == "number" and not _ID_LIKE.search(fold(c))]
    dates = [c for c in t.columns if t.types.get(c) == "date"]
    cats = []
    for c in t.columns:
        if t.types.get(c) != "text":
            continue
        distinct = {v for v in t.values(c) if v is not None}
        if 2 <= len(distinct) <= 60 and len(distinct) < max(3, len(t.rows) * 0.9):
            cats.append((len(distinct), c))
    cats = [c for _n, c in sorted(cats)]
    if not cats:
        # A short list where every name is different (5 products, 30 pupils) still reads best
        # as a bar per name (live 2026-10-24: only a scatter of two numbers was made).
        for c in t.columns:
            distinct = {v for v in t.values(c) if v is not None}
            if t.types.get(c) == "text" and 2 <= len(distinct) <= 60:
                cats.append(c)
                break
    return {"numbers": nums, "dates": dates, "categories": cats}


def auto_dashboard(
    t: Table, source: str, name: str = "", sheet: str | None = None
) -> dict[str, Any]:
    """A sensible first dashboard; the user edits it afterwards."""
    t = typed(t) if not t.types else t
    p = profile(t)
    steps: list[dict[str, Any]] = []
    rev, exp = _find(p["numbers"], _REVENUE), _find(p["numbers"], _EXPENSE)
    measures = list(p["numbers"])
    if rev and exp and rev != exp:
        steps.append({"op": "calc", "name": "Sof foyda", "expr": f"[{rev}] - [{exp}]"})
        measures = [rev, exp, "Sof foyda"] + [m for m in measures if m not in (rev, exp)]
    elif total := _find(measures, ("jami", "summa", "total", "itogo", "umumiy")):
        measures = [total] + [m for m in measures if m != total]  # the amount, not «Soni»
    main = measures[0] if measures else None
    date = p["dates"][0] if p["dates"] else None
    cat = p["categories"][0] if p["categories"] else None
    cat2 = p["categories"][1] if len(p["categories"]) > 1 else None
    n = 0

    def wid() -> str:
        nonlocal n
        n += 1
        return f"w{n}"

    overview: list[dict[str, Any]] = []
    by_day = False
    for m in measures[:3]:
        overview.append({"id": wid(), "type": "kpi", "title": f"{m} ({AGG_LABELS['sum']})",
                         "y": m, "agg": "sum", "growth": bool(date), "x": date})  # fmt: skip
    if rev and exp:
        overview.append({"id": wid(), "type": "kpi", "title": "Rentabellik, %", "y": "Sof foyda",
                         "agg": "margin", "base": rev})  # fmt: skip
    overview.append({"id": wid(), "type": "kpi", "title": "Qatorlar soni", "agg": "count"})
    if date and main:
        # All dates in one month: a month line is a single point (live 2026-10-24: an empty
        # chart on the phone) — show the days instead.
        months = {(d.year, d.month) for d in t.values(date) if hasattr(d, "month")}
        by_day = len(months) <= 1
        overview.append({"id": wid(), "type": "line",
                         "title": "Kunlar bo‘yicha" if by_day else "Oylar bo‘yicha", "x": date,
                         "y": measures[:3], "agg": "sum",
                         "period": "day" if by_day else "month"})  # fmt: skip
    if cat and main:
        overview.append({"id": wid(), "type": "bar", "title": f"{cat} bo‘yicha {main}", "x": cat,
                         "y": main, "agg": "sum", "top": 12})  # fmt: skip
        overview.append({"id": wid(), "type": "pie", "title": f"{main} ulushi ({cat2 or cat})",
                         "x": cat2 or cat, "y": main, "agg": "sum", "top": 6})  # fmt: skip
    if cat:
        overview.append({"id": wid(), "type": "table", "title": f"{cat} kesimida", "x": cat,
                         "y": measures[:4], "agg": "sum"})  # fmt: skip
    details: list[dict[str, Any]] = []
    if len(measures) >= 2:
        details.append({"id": wid(), "type": "scatter", "title": f"{measures[0]} va {measures[1]}",
                        "x": measures[0], "y": measures[1]})  # fmt: skip
    if cat2 and main:
        details.append({"id": wid(), "type": "bar", "title": f"{cat2} bo‘yicha {main}", "x": cat2,
                        "y": main, "agg": "sum", "top": 15})  # fmt: skip
    if date and main and not (rev and exp) and not by_day:  # by_day: already on «Umumiy»
        details.append({"id": wid(), "type": "line", "title": "Kunlar bo‘yicha", "x": date,
                        "y": [main], "agg": "sum", "period": "day"})  # fmt: skip
    details.append({"id": wid(), "type": "table", "title": "Barcha ma’lumotlar"})
    pages = [{"title": "Umumiy", "widgets": overview}, {"title": "Batafsil", "widgets": details}]
    return {
        "name": name or Path(source).stem,
        "source": source,
        "sheet": sheet,
        "steps": steps,
        "slicers": p["categories"][:3],
        "date_col": date,
        "pages": pages,
    }


# ---------------- computing widgets ----------------


def _agg(values: list[float], how: str) -> float | None:
    vals = [v for v in values if v is not None]
    if how == "count":
        return float(len(values))
    if not vals:
        return None
    if how == "sum":
        return float(sum(vals))
    if how == "avg":
        return sum(vals) / len(vals)
    if how == "min":
        return float(min(vals))
    if how == "max":
        return float(max(vals))
    raise ValueError(f"noma’lum hisob: {how}")


def period_key(d: dt.date, period: str) -> str:
    if period == "day":
        return d.isoformat()
    if period == "year":
        return str(d.year)
    if period == "quarter":
        return f"{d.year}-Q{(d.month - 1) // 3 + 1}"
    return f"{d.year}-{d.month:02d}"


def filtered(t: Table, state: dict[str, Any], date_col: str | None) -> Table:
    """Slicer selections ({col: [values]}) and the date range applied."""
    out = t.copy()
    for col, wanted in (state.get("filters") or {}).items():
        if not wanted or col not in out.columns:
            continue
        j = out.col(col)
        keep = {str(w) for w in wanted}
        out.rows = [r for r in out.rows if _label(r[j]) in keep]
    if date_col and date_col in out.columns:
        j = out.col(date_col)
        lo = _date(state.get("date_from"))
        hi = _date(state.get("date_to"))
        if lo or hi:
            out.rows = [r for r in out.rows if r[j] is not None and (not lo or r[j] >= lo)
                        and (not hi or r[j] <= hi)]  # fmt: skip
    return out


def _date(v: Any) -> dt.date | None:
    if not v:
        return None
    try:
        return dt.date.fromisoformat(str(v)[:10])
    except ValueError:
        return None


def _label(v: Any) -> str:
    if v is None:
        return "(bo‘sh)"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.date):
        return v.isoformat()
    return str(v)


def _ys(w: dict[str, Any]) -> list[str]:
    y = w.get("y")
    if not y:
        return []
    return list(y) if isinstance(y, list) else [y]


def _group(t: Table, x: str, ys: list[str], agg: str, period: str | None = None):
    j = t.col(x)
    idx = [t.col(y) for y in ys]
    groups: dict[str, list[list[Any]]] = defaultdict(list)
    for r in t.rows:
        key = r[j]
        if period:
            if key is None:
                continue
            key = period_key(key, period)
        groups[_label(key) if not period else key].append(r)
    return {k: [_agg([r[i] for r in rows], agg) for i in idx] for k, rows in groups.items()}


def compute_widget(t: Table, w: dict[str, Any], date_col: str | None) -> dict[str, Any]:
    kind = w.get("type")
    out = {k: w.get(k) for k in ("id", "type", "title", "x", "y", "agg", "period")}
    ys = _ys(w)
    agg = w.get("agg") or "sum"
    if kind == "kpi":
        if agg == "count":
            out["value"] = float(len(t.rows))
            return out
        if agg == "margin":  # profit / revenue, %
            p, base = _agg(t.values(ys[0]), "sum"), _agg(t.values(w["base"]), "sum")
            out["value"] = (p / base * 100) if p is not None and base else None
            out["unit"] = "%"
            return out
        out["value"] = _agg(t.values(ys[0]), agg)
        dcol = w.get("x") or date_col
        if w.get("growth") and dcol and t.types.get(dcol) == "date":
            series = _group(t, dcol, ys[:1], agg, "month")
            keys = sorted(series)
            if len(keys) >= 2:
                last, prev = series[keys[-1]][0], series[keys[-2]][0]
                out["growth"] = (
                    ((last - prev) / abs(prev) * 100) if last is not None and prev else None
                )
                out["growth_label"] = f"{keys[-1]} / {keys[-2]}"
        return out
    if kind in ("bar", "pie"):
        x = w["x"]
        groups = _group(t, x, ys[:1] or [x], agg if ys else "count")
        items = sorted(
            ((k, v[0]) for k, v in groups.items() if v[0] is not None), key=lambda kv: -kv[1]
        )
        top = int(w.get("top") or 12)
        if len(items) > top:
            rest = sum(v for _k, v in items[top - 1 :])
            items = items[: top - 1] + [("Boshqalar", rest)]
        out["labels"] = [k for k, _v in items]
        out["values"] = [v for _k, v in items]
        return out
    if kind == "line":
        x = w["x"]
        period = w.get("period") or "month"
        is_date = t.types.get(x) == "date"
        groups = _group(t, x, ys, agg, period if is_date else None)
        keys = sorted(groups)
        out["labels"] = keys
        out["series"] = [
            {"name": y, "values": [groups[k][i] for k in keys]} for i, y in enumerate(ys)
        ]
        return out
    if kind == "scatter":
        xi, yi = t.col(w["x"]), t.col(ys[0])
        pts = [(r[xi], r[yi]) for r in t.rows if r[xi] is not None and r[yi] is not None]
        step = max(1, len(pts) // 1500)
        out["points"] = pts[::step]
        return out
    if kind == "table":
        if w.get("x"):
            groups = _group(t, w["x"], ys, agg)
            rows = sorted(groups.items(), key=lambda kv: -(kv[1][0] or 0) if ys else kv[0])
            out["columns"] = [w["x"]] + [f"{y} ({AGG_LABELS.get(agg, agg)})" for y in ys]
            out["rows"] = [[k, *v] for k, v in rows[:300]]
        else:
            out["columns"] = list(t.columns)
            out["rows"] = [[_cell(v) for v in r] for r in t.rows[:300]]
        out["total_rows"] = len(t.rows)
        return out
    raise ValueError(f"noma’lum vidjet: {kind}")


def _cell(v: Any) -> Any:
    return v.isoformat() if isinstance(v, dt.date) else v


def prepare(cfg: dict[str, Any]) -> Table:
    """Load the source and run the ETL steps."""
    t = typed(load(cfg["source"], cfg.get("sheet")))
    return apply_steps(t, cfg.get("steps") or [])


def compute(cfg: dict[str, Any], state: dict[str, Any] | None = None, table: Table | None = None):
    state = state or {}
    base = table if table is not None else prepare(cfg)
    t = filtered(base, state, cfg.get("date_col"))
    page = min(int(state.get("page") or 0), max(0, len(cfg["pages"]) - 1))
    widgets, problems = [], []
    for w in cfg["pages"][page]["widgets"] if cfg["pages"] else []:
        try:
            widgets.append(compute_widget(t, w, cfg.get("date_col")))
        except (LookupError, ValueError, TypeError, KeyError) as e:
            problems.append(f"{w.get('title') or w.get('type')}: {e}")
    slicers = {}
    for c in cfg.get("slicers") or []:
        if c in base.columns:
            vals = sorted({_label(v) for v in base.values(c)})
            slicers[c] = vals[:300]
    dates = []
    dc = cfg.get("date_col")
    if dc and dc in base.columns:
        dates = [v for v in base.values(dc) if v is not None]
    return {
        "page": page,
        "pages": [p["title"] for p in cfg["pages"]],
        "widgets": widgets,
        "problems": problems,
        "slicers": slicers,
        "date_range": [min(dates).isoformat(), max(dates).isoformat()] if dates else None,
        "rows": len(t.rows),
        "total_rows": len(base.rows),
        "columns": [{"name": c, "type": base.types.get(c, "text")} for c in base.columns],
        "notes": base.notes,
    }


# ---------------- storage ----------------


def bi_dir() -> Path:
    return app_config.data_dir() / "BI"


def _slug(name: str) -> str:
    s = re.sub(r"[^\w‘’' -]", "", name).strip().replace(" ", "_")
    return s[:60] or "dashboard"


def save(cfg: dict[str, Any]) -> Path:
    validate(cfg)
    d = bi_dir()
    d.mkdir(parents=True, exist_ok=True)
    path = d / f"{_slug(cfg['name'])}.json"
    cfg = {**cfg, "saved": time.time()}
    path.write_text(json.dumps(cfg, ensure_ascii=False, indent=1), encoding="utf-8")
    return path


def validate(cfg: dict[str, Any]) -> None:
    if not str(cfg.get("name") or "").strip():
        raise ValueError("Dashboard nomini yozing")
    if not cfg.get("source"):
        raise ValueError("Ma’lumot manbasi tanlanmagan")
    for p in cfg.get("pages") or []:
        for w in p.get("widgets") or []:
            if w.get("type") not in WIDGET_TYPES:
                raise ValueError(f"noma’lum vidjet turi: {w.get('type')}")
            if w.get("agg") and w["agg"] not in (*AGGS, "margin"):
                raise ValueError(f"noma’lum hisob: {w['agg']}")


def saved() -> list[dict[str, Any]]:
    d = bi_dir()
    if not d.is_dir():
        return []
    out = []
    for p in sorted(d.glob("*.json")):
        try:
            cfg = json.loads(p.read_text(encoding="utf-8"))
        except ValueError:
            continue
        out.append({"file": str(p), "name": cfg.get("name"), "source": cfg.get("source"),
                    "saved": cfg.get("saved")})  # fmt: skip
    return sorted(out, key=lambda x: -(x["saved"] or 0))


def load_saved(file: str) -> dict[str, Any]:
    p = Path(file)
    if bi_dir().resolve() not in p.resolve().parents:
        raise ValueError("bu NEXVARA BI fayli emas")
    return json.loads(p.read_text(encoding="utf-8"))


# ---------------- export ----------------


def export_rows(cfg: dict[str, Any], state: dict[str, Any], fmt: str, out: Path) -> Path:
    """The filtered data as CSV or XLSX."""
    t = filtered(prepare(cfg), state, cfg.get("date_col"))
    out.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "csv":
        import csv

        with out.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(t.columns)
            for r in t.rows:
                w.writerow(["" if v is None else _cell(v) for v in r])
        return out
    if fmt == "xlsx":
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ma’lumot"[:31]
        ws.append(t.columns)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="6366F1")
        for r in t.rows:
            ws.append(list(r))
        for j, c in enumerate(t.columns, 1):
            if t.types.get(c) == "number":
                for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    # Stored in Excel's invariant form; Excel shows the user's separators.
                    row[0].number_format = "#,##0.00" if any(
                        isinstance(v, float) and not v.is_integer() for v in t.values(c) if v
                    ) else "#,##0"  # fmt: skip
            elif t.types.get(c) == "date":
                for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    row[0].number_format = "DD.MM.YYYY"
            width = max([len(str(c))] + [len(str(v)) for v in t.values(c)[:500] if v is not None])
            ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = min(45, width + 3)
        ws.freeze_panes = "A2"
        wb.save(out)
        return out
    raise ValueError("format: csv yoki xlsx")


def save_png(data_url: str, out: Path) -> Path:
    import base64

    head, _, b64 = data_url.partition(",")
    if not head.startswith("data:image/png;base64"):
        raise ValueError("PNG rasm kutilgan edi")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(base64.b64decode(b64))
    return out


def pngs_to_pdf(data_urls: list[str], out: Path, title: str = "") -> Path:
    """One landscape A4 page per dashboard page image."""
    import base64

    import pymupdf

    doc = pymupdf.open()
    for url in data_urls:
        pix = pymupdf.Pixmap(base64.b64decode(url.partition(",")[2]))
        if pix.alpha:
            pix = pymupdf.Pixmap(pix, 0)
        jpg = pix.tobytes("jpeg", jpg_quality=90)  # a raw image would make a 10 MB page
        page = doc.new_page(width=842, height=595)
        if title:
            page.insert_text((36, 30), title, fontsize=13, fontname="helv")
        page.insert_image(pymupdf.Rect(36, 40, 806, 575), stream=jpg, keep_proportion=True)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    doc.close()
    return out


def export_dir() -> Path:
    return app_config.data_dir() / "BI" / "Eksport"
