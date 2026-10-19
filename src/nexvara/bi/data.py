"""Loading tables and the ETL steps (types, cleaning, joins, calculated columns, filters).

A table is a list of column names and rows of Python values. After `typed()` every
column has one type: "number" (float), "date" (datetime.date) or "text" (str); empty and
"-", "n/a" become None.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

SUPPORTED = (".xlsx", ".xlsm", ".csv", ".json", ".xml", ".txt", ".tsv")
DATABASES = (".db", ".sqlite", ".sqlite3", ".accdb", ".mdb")  # bi/db.py, read-only
MAX_ROWS = 200_000
NULLS = {"", "-", "—", "n/a", "na", "null", "none", "nan", "#n/a", "yo‘q", "yoq"}


@dataclass
class Table:
    columns: list[str]
    rows: list[list[Any]]
    types: dict[str, str] = field(default_factory=dict)
    source: str = ""
    notes: list[str] = field(default_factory=list)

    def col(self, name: str) -> int:
        try:
            return self.columns.index(name)
        except ValueError:
            raise LookupError(f"«{name}» ustuni topilmadi") from None

    def values(self, name: str) -> list[Any]:
        i = self.col(name)
        return [r[i] for r in self.rows]

    def copy(self) -> Table:
        return Table(list(self.columns), [list(r) for r in self.rows], dict(self.types),
                     self.source, list(self.notes))  # fmt: skip


# ---------------- loading ----------------


def _clean_header(cols: list[Any]) -> list[str]:
    out: list[str] = []
    for i, c in enumerate(cols, 1):
        name = str(c).strip() if c not in (None, "") else f"Ustun {i}"
        base, n = name, 2
        while name in out:  # duplicate headers get a number
            name = f"{base} ({n})"
            n += 1
        out.append(name)
    return out


def _from_grid(grid: list[list[Any]], source: str) -> Table:
    grid = [r for r in grid if any(v not in (None, "") for v in r)]
    if not grid:
        raise ValueError("Faylda ma’lumot yo‘q")
    width = max(len(r) for r in grid)
    grid = [list(r) + [None] * (width - len(r)) for r in grid]
    # Drop columns that are empty everywhere (Excel formatting leftovers).
    keep = [j for j in range(width) if any(r[j] not in (None, "") for r in grid)]
    grid = [[r[j] for j in keep] for r in grid]
    header, *rows = grid
    t = Table(_clean_header(header), rows[:MAX_ROWS], source=source)
    if len(rows) > MAX_ROWS:
        t.notes.append(f"Faqat birinchi {MAX_ROWS} qator olindi ({len(rows)} ta bor edi)")
    return t


def sheets(path: str | Path) -> list[str]:
    """Sheets of a workbook, or tables/views of a database."""
    p = Path(path)
    if p.suffix.lower() in DATABASES:
        from nexvara.bi import db

        return db.tables(p)
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        return []
    import openpyxl

    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _xlsx(p: Path, sheet: str | None) -> Table:
    import openpyxl

    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    # Header = first row with at least two filled cells (titles above a table are skipped).
    start = next((i for i, r in enumerate(grid) if sum(v not in (None, "") for v in r) >= 2), 0)
    return _from_grid(grid[start:], str(p))


def _read_text(p: Path) -> str:
    raw = p.read_bytes()
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _delimited(p: Path) -> Table:
    text = _read_text(p)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = max(";,\t|", key=sample.count)
    grid = list(csv.reader(io.StringIO(text), delimiter=delim))
    return _from_grid(grid, str(p))


def _records_table(records: list[dict], source: str) -> Table:
    cols: list[str] = []
    for r in records:
        for k in r:
            if k not in cols:
                cols.append(k)
    rows = [[_scalar(r.get(c)) for c in cols] for r in records]
    return _from_grid([cols, *rows], source)


def _scalar(v: Any) -> Any:
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return v


def _json(p: Path) -> Table:
    data = json.loads(_read_text(p))
    if isinstance(data, dict):
        # {"data": [...]} or {"items": [...]}: the first list of objects inside.
        data = next((v for v in data.values() if isinstance(v, list)), [data])
    if not isinstance(data, list) or not data:
        raise ValueError("JSON ichida jadval (ro‘yxat) topilmadi")
    if all(isinstance(r, dict) for r in data):
        return _records_table(data, str(p))
    if all(isinstance(r, list) for r in data):
        return _from_grid(data, str(p))
    raise ValueError("JSON ro‘yxati bir xil tuzilishda emas")


def _xml(p: Path) -> Table:
    root = ET.fromstring(p.read_bytes())
    # The repeated element: the child tag that occurs most often (one level down or two).
    best: list[ET.Element] = []
    for parent in [root, *list(root)]:
        tags: dict[str, list[ET.Element]] = {}
        for ch in parent:
            tags.setdefault(ch.tag, []).append(ch)
        for items in tags.values():
            if len(items) > len(best):
                best = items
    if not best:
        raise ValueError("XML ichida takrorlanuvchi yozuvlar topilmadi")
    records = []
    for el in best:
        rec = dict(el.attrib)
        for ch in el:
            rec[ch.tag] = (ch.text or "").strip()
        if not len(el) and (el.text or "").strip():
            rec[el.tag] = el.text.strip()
        records.append(rec)
    return _records_table(records, str(p))


def load(path: str | Path, sheet: str | None = None) -> Table:
    p = Path(path)
    if p.is_dir():
        return load_folder(p)
    if not p.is_file():
        raise FileNotFoundError(f"Fayl topilmadi: {p}")
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return _xlsx(p, sheet)
    if ext in (".csv", ".txt", ".tsv"):
        return _delimited(p)
    if ext == ".json":
        return _json(p)
    if ext == ".xml":
        return _xml(p)
    if ext in DATABASES:  # sheet = table name or one SELECT query
        from nexvara.bi import db

        cols, rows = db.read(p, sheet, MAX_ROWS)
        t = Table(_clean_header(cols), rows[:MAX_ROWS], source=str(p))
        if len(rows) > MAX_ROWS:
            t.notes.append(f"Faqat birinchi {MAX_ROWS} qator olindi")
        if not rows:
            t.notes.append("Jadval bo‘sh")
        return t
    if ext == ".xls":
        raise ValueError("Eski .xls formatini Excel'da .xlsx qilib saqlang (Fayllar → PDF emas)")
    raise ValueError(f"{ext} fayl turi qo‘llanmaydi. Mumkin: {', '.join(SUPPORTED + DATABASES)}")


def load_folder(folder: Path) -> Table:
    """All supported files of a folder stacked into one table, with a "Fayl" column."""
    files = sorted(f for f in folder.iterdir() if f.is_file() and f.suffix.lower() in SUPPORTED
                   and not f.name.startswith("~$"))  # fmt: skip
    if not files:
        raise ValueError("Papkada mos fayl yo‘q (xlsx, csv, json, xml, txt)")
    tables = [load(f) for f in files]
    cols: list[str] = ["Fayl"]
    for t in tables:
        cols += [c for c in t.columns if c not in cols]
    rows = []
    for f, t in zip(files, tables, strict=True):
        idx = {c: i for i, c in enumerate(t.columns)}
        rows += [[f.name] + [r[idx[c]] if c in idx else None for c in cols[1:]] for r in t.rows]
    return Table(cols, rows[:MAX_ROWS], source=str(folder))


# ---------------- types ----------------

_NUM = re.compile(r"^[-+]?(?:\d[\d\s  '.,]*\d|\d)$|^[-+]?[.,]\d+$")
_DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y.%m.%d", "%d-%m-%Y", "%Y/%m/%d",
                 "%d.%m.%y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S",
                 "%d.%m.%Y %H:%M")  # fmt: skip
_MONTHS = {"yanvar": 1, "fevral": 2, "mart": 3, "aprel": 4, "may": 5, "iyun": 6, "iyul": 7,
           "avgust": 8, "sentabr": 9, "oktabr": 10, "noyabr": 11, "dekabr": 12}  # fmt: skip


def parse_number(v: Any) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if not isinstance(v, str):
        return None
    s = v.strip().replace("%", "")
    for cur in ("so‘m", "so'm", "som", "sum", "$", "usd", "uzs", "€", "₽"):
        s = s.lower().replace(cur, "").strip()
    if not s or not _NUM.match(s):
        return None
    s = re.sub(r"[\s  ']", "", s)
    if "," in s and "." in s:  # 1.234,5 or 1,234.5: the last one is the decimal mark
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        s = s.replace("." if dec == "," else ",", "").replace(",", ".")
    elif s.count(",") == 1:  # Uzbek/Russian style: "1 234,5"
        s = s.replace(",", ".")
    elif s.count(",") > 1:  # "1,234,567"
        s = s.replace(",", "")
    elif s.count(".") > 1:  # "1.234.567"
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def parse_date(v: Any) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if not isinstance(v, str):
        return None
    s = v.strip()
    for f in _DATE_FORMATS:
        try:
            return dt.datetime.strptime(s, f).date()
        except ValueError:
            continue
    m = re.match(r"^(\d{4})[-./](\d{1,2})$", s)  # 2026-03 (a month)
    if m and 1 <= int(m.group(2)) <= 12:
        return dt.date(int(m.group(1)), int(m.group(2)), 1)
    low = s.lower()
    for name, mon in _MONTHS.items():  # "Mart 2026", "2026 mart"
        if name in low:
            y = re.search(r"\d{4}", low)
            if y:
                return dt.date(int(y.group(0)), mon, 1)
    return None


def is_null(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip().lower() in NULLS)


def detect_type(values: list[Any]) -> str:
    vals = [v for v in values if not is_null(v)]
    if not vals:
        return "text"
    sample = vals[:2000]
    dates = sum(parse_date(v) is not None for v in sample)
    nums = sum(parse_number(v) is not None for v in sample)
    if dates >= 0.9 * len(sample) and not all(isinstance(v, (int, float)) for v in sample):
        return "date"
    if nums >= 0.9 * len(sample):
        return "number"
    return "text"


def typed(t: Table) -> Table:
    """Detect each column's type and convert values (unparseable ones become None)."""
    out = t.copy()
    for j, c in enumerate(out.columns):
        kind = detect_type([r[j] for r in out.rows])
        out.types[c] = kind
        conv = {"number": parse_number, "date": parse_date}.get(kind)
        for r in out.rows:
            v = r[j]
            if is_null(v):
                r[j] = None
            elif conv:
                r[j] = conv(v)
            else:
                r[j] = str(v).strip() if not isinstance(v, str) else re.sub(r"\s+", " ", v).strip()
    return out


# ---------------- ETL steps ----------------

STEP_LABELS = {
    "dedupe": "Takroriy qatorlarni o‘chirish",
    "drop_empty": "Bo‘sh qatorlarni o‘chirish",
    "fill_null": "Bo‘sh qiymatlarni to‘ldirish",
    "drop_null": "Ustunda bo‘sh bo‘lgan qatorlarni o‘chirish",
    "filter": "Filtr",
    "calc": "Hisoblangan ustun",
    "rename": "Ustun nomini o‘zgartirish",
    "drop_col": "Ustunni olib tashlash",
    "join": "Boshqa fayl bilan birlashtirish",
    "set_type": "Ustun turini o‘zgartirish",
    "append": "Boshqa fayl qatorlarini qo‘shish",
}
_OPS = {"=", "!=", ">", ">=", "<", "<=", "contains"}
_EXPR_TOKEN = re.compile(r"\s*(?:(\d+(?:[.,]\d+)?)|\[([^\]]+)\]|([-+*/()]))")


def describe_step(s: dict[str, Any]) -> str:
    op = s.get("op")
    label = STEP_LABELS.get(op, op)
    if op == "calc":
        return f"{label}: {s.get('name')} = {s.get('expr')}"
    if op == "filter":
        return f"{label}: {s.get('col')} {s.get('cmp')} {s.get('value')}"
    if op in ("fill_null", "drop_null", "drop_col", "set_type"):
        extra = f" → {s.get('value', s.get('type', ''))}" if op in ("fill_null", "set_type") else ""
        return f"{label}: {s.get('col')}{extra}"
    if op == "rename":
        return f"{label}: {s.get('col')} → {s.get('name')}"
    if op == "join":
        return f"{label}: {Path(str(s.get('file'))).name} ({s.get('on')})"
    if op == "append":
        return f"{label}: {Path(str(s.get('file'))).name}"
    return label


def eval_expr(expr: str, row: dict[str, Any]) -> float | None:
    """'[Daromad] - [Xarajat]', '[Narx] * [Soni] * 1,12'. Only numbers, columns in [],
    + - * / and brackets: no Python eval."""
    pos, tokens = 0, []
    while pos < len(expr):
        m = _EXPR_TOKEN.match(expr, pos)
        if not m or m.end() == pos:
            if expr[pos:].strip():
                raise ValueError(f"Formulani tushunmadim: «{expr[pos:].strip()}»")
            break
        pos = m.end()
        if m.group(1):
            tokens.append(float(m.group(1).replace(",", ".")))
        elif m.group(2):
            name = m.group(2).strip()
            if name not in row:
                raise LookupError(f"Formulada «{name}» ustuni yo‘q")
            tokens.append(("col", name))
        else:
            tokens.append(m.group(3))
    vals = [row[t[1]] if isinstance(t, tuple) else t for t in tokens]
    if any(v is None for v in vals):
        return None
    it = iter(vals)
    cur = [next(it, None)]

    def parse_sum():
        v = parse_prod()
        while cur[0] in ("+", "-"):
            op = cur[0]
            cur[0] = next(it, None)
            w = parse_prod()
            v = v + w if op == "+" else v - w
        return v

    def parse_prod():
        v = parse_atom()
        while cur[0] in ("*", "/"):
            op = cur[0]
            cur[0] = next(it, None)
            w = parse_atom()
            if op == "*":
                v = v * w
            else:
                v = v / w if w else None
                if v is None:
                    raise ZeroDivisionError
        return v

    def parse_atom():
        t = cur[0]
        cur[0] = next(it, None)
        if t == "-":
            return -parse_atom()
        if t == "(":
            v = parse_sum()
            if cur[0] != ")":
                raise ValueError("Qavs yopilmagan")
            cur[0] = next(it, None)
            return v
        if isinstance(t, (int, float)):
            return float(t)
        raise ValueError("Formula noto‘g‘ri")

    try:
        v = parse_sum()
    except ZeroDivisionError:
        return None
    if cur[0] is not None:
        raise ValueError("Formula noto‘g‘ri")
    return v


def _match(v: Any, cmp: str, target: Any, kind: str) -> bool:
    if v is None:
        return False
    if kind == "number":
        target = parse_number(target)
    elif kind == "date":
        target = parse_date(target)
    if target is None:
        return False
    if cmp == "contains":
        return str(target).lower() in str(v).lower()
    if kind == "text":
        v, target = str(v).lower(), str(target).lower()
    return {
        "=": v == target, "!=": v != target, ">": v > target, ">=": v >= target,
        "<": v < target, "<=": v <= target,
    }[cmp]  # fmt: skip


def apply_steps(t: Table, steps: list[dict[str, Any]], loader=None) -> Table:
    """Run the saved transformation steps in order on a typed table."""
    loader = loader or (lambda p: typed(load(p)))
    t = t.copy()
    for s in steps:
        op = s.get("op")
        if op == "dedupe":
            seen, rows = set(), []
            for r in t.rows:
                key = tuple(r)
                if key not in seen:
                    seen.add(key)
                    rows.append(r)
            t.notes.append(f"{len(t.rows) - len(rows)} ta takroriy qator o‘chirildi")
            t.rows = rows
        elif op == "drop_empty":
            t.rows = [r for r in t.rows if any(v is not None for v in r)]
        elif op == "drop_null":
            j = t.col(s["col"])
            t.rows = [r for r in t.rows if r[j] is not None]
        elif op == "fill_null":
            j = t.col(s["col"])
            kind = t.types.get(s["col"], "text")
            fill = s.get("value")
            if fill == "__zero__" or (kind == "number" and fill in (None, "")):
                fill = 0.0
            elif fill == "__mean__":
                nums = [r[j] for r in t.rows if r[j] is not None]
                fill = sum(nums) / len(nums) if nums else 0.0
            elif kind == "number":
                fill = parse_number(fill)
            elif kind == "date":
                fill = parse_date(fill)
            for r in t.rows:
                if r[j] is None:
                    r[j] = fill
        elif op == "filter":
            if s.get("cmp") not in _OPS:
                raise ValueError("noma’lum taqqoslash")
            j = t.col(s["col"])
            kind = t.types.get(s["col"], "text")
            t.rows = [r for r in t.rows if _match(r[j], s["cmp"], s.get("value"), kind)]
        elif op == "calc":
            name = str(s.get("name") or "").strip()
            if not name:
                raise ValueError("Yangi ustun nomini yozing")
            vals = [eval_expr(str(s.get("expr") or ""), dict(zip(t.columns, r, strict=True)))
                    for r in t.rows]  # fmt: skip
            if name in t.columns:
                j = t.col(name)
                for r, v in zip(t.rows, vals, strict=True):
                    r[j] = v
            else:
                t.columns.append(name)
                for r, v in zip(t.rows, vals, strict=True):
                    r.append(v)
            t.types[name] = "number"
        elif op == "rename":
            j = t.col(s["col"])
            new = str(s.get("name") or "").strip()
            if not new or new in t.columns:
                raise ValueError("Yangi nom bo‘sh yoki band")
            t.types[new] = t.types.pop(s["col"], "text")
            t.columns[j] = new
        elif op == "drop_col":
            j = t.col(s["col"])
            t.columns.pop(j)
            t.types.pop(s["col"], None)
            for r in t.rows:
                r.pop(j)
        elif op == "set_type":
            j = t.col(s["col"])
            kind = s.get("type")
            conv = {
                "number": parse_number,
                "date": parse_date,
                "text": lambda v: None if v is None else str(v),
            }
            if kind not in conv:
                raise ValueError("tur: number, date yoki text")
            for r in t.rows:
                r[j] = conv[kind](r[j] if not isinstance(r[j], dt.date) else r[j].isoformat())
            t.types[s["col"]] = kind
        elif op == "join":
            other = loader(s["file"])
            on = s["on"]
            j, k = t.col(on), other.col(on)
            index: dict[Any, list[Any]] = {}
            for r in other.rows:
                index.setdefault(_key(r[k]), r)
            extra = [c for c in other.columns if c != on]
            names = [
                c if c not in t.columns else f"{c} ({Path(str(s['file'])).stem})" for c in extra
            ]
            idx = [other.col(c) for c in extra]
            for r in t.rows:
                match = index.get(_key(r[j]))
                r += [match[i] if match else None for i in idx]
            for c, n in zip(extra, names, strict=True):
                t.columns.append(n)
                t.types[n] = other.types.get(c, "text")
            matched = sum(_key(r[j]) in index for r in t.rows)
            t.notes.append(f"Birlashtirildi: {matched}/{len(t.rows)} qator mos keldi")
        elif op == "append":
            # Rows of another file below these ones (same-named columns line up; new
            # columns are added; "Manba" tells which file a row came from).
            other = loader(s["file"])
            if "Manba" not in t.columns:
                t.columns.append("Manba")
                t.types["Manba"] = "text"
                first = Path(str(t.source or "")).name or "asosiy"
                for r in t.rows:
                    r.append(first)
            for c in other.columns:
                if c not in t.columns:
                    t.columns.append(c)
                    t.types[c] = other.types.get(c, "text")
                    for r in t.rows:
                        r.append(None)
            idx = {c: i for i, c in enumerate(other.columns)}
            name = Path(str(s["file"])).name
            added = 0
            for r in other.rows:
                row = [r[idx[c]] if c in idx else None for c in t.columns]
                row[t.col("Manba")] = name
                t.rows.append(row)
                added += 1
            t.rows = t.rows[:MAX_ROWS]
            t.notes.append(f"{name} dan {added} ta qator qo‘shildi")
        else:
            raise ValueError(f"noma’lum bosqich: {op}")
    return t


def _key(v: Any) -> Any:
    return str(v).strip().lower() if isinstance(v, str) else v
