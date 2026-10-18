"""Answer sheets ("javob varaqasi"): print a sheet, photograph the filled one, get the score.

The sheet has four solid corner squares. On a photo (taken at an angle, rotated, any phone)
they are found as dark square blobs; a homography maps the sheet's own coordinates to the
photo, so every bubble is read where it really is. Pure Python + PyMuPDF (no OpenCV, no AI,
nothing leaves the PC). Results go to an Excel file per sheet set.
"""

from __future__ import annotations

import json
import time
import uuid
from pathlib import Path
from typing import Any

from nexvara import config
from nexvara.core.task import UserError

W, H = 595.0, 842.0  # A4 in points
MARK = 26.0  # corner square side
MARKS = [
    (30 + MARK / 2, 30 + MARK / 2),
    (W - 30 - MARK / 2, 30 + MARK / 2),
    (30 + MARK / 2, H - 30 - MARK / 2),
    (W - 30 - MARK / 2, H - 30 - MARK / 2),
]
LETTERS = "ABCD"
MAX_Q = 60
R = 8.5  # bubble radius
ROWS = 20
COLS_X = (70.0, 255.0, 440.0)
TOP = 210.0
ROW_H = 26.0


def folder() -> Path:
    p = config.data_dir() / "Javob varaqalari"
    p.mkdir(parents=True, exist_ok=True)
    return p


def layout(n: int, variants: int) -> dict[str, Any]:
    """Bubble centres in sheet points."""
    if not 1 <= n <= MAX_Q:
        raise UserError(f"Savollar soni 1–{MAX_Q} bo‘lsin")
    if not 1 <= variants <= 4:
        raise UserError("Variantlar 1–4 bo‘lsin")
    q = []
    for i in range(n):
        x0 = COLS_X[i // ROWS]
        y = TOP + (i % ROWS) * ROW_H
        q.append([(x0 + 34 + k * 28, y) for k in range(4)])
    var = [(170 + k * 34, 158.0) for k in range(variants)]
    return {"q": q, "variant": var}


def make_sheet_pdf(path: Path, title: str, n: int, variants: int = 1, copies: int = 1) -> Path:
    import fitz

    lay = layout(n, variants)
    doc = fitz.open()
    for _ in range(max(1, min(copies, 60))):
        page = doc.new_page(width=W, height=H)
        for cx, cy in MARKS:
            page.draw_rect(
                fitz.Rect(cx - MARK / 2, cy - MARK / 2, cx + MARK / 2, cy + MARK / 2),
                color=(0, 0, 0),
                fill=(0, 0, 0),
            )
        safe = title.replace("‘", "'").replace("’", "'")[:70]
        page.insert_text((70, 75), safe, fontsize=15, fontname="helv")
        page.insert_text(
            (70, 105),
            "F.I.Sh: ________________________________   Sinf: ______",
            fontsize=11,
            fontname="helv",
        )
        page.insert_text(
            (70, 128),
            "Sana: ____________   Doirachani ruchka bilan to'liq bo'yang.",
            fontsize=9,
            fontname="helv",
        )
        page.insert_text((70, 162), "Variant:", fontsize=11, fontname="helv")
        for k, (x, y) in enumerate(lay["variant"]):
            page.draw_circle((x, y), R, color=(0, 0, 0), width=0.8)
            page.insert_text((x - 3, y - R - 3), LETTERS[k], fontsize=8, fontname="helv")
        cols = (n - 1) // ROWS + 1
        for c in range(cols):
            for k, ch in enumerate(LETTERS):
                page.insert_text(
                    (COLS_X[c] + 34 + k * 28 - 3, TOP - 16), ch, fontsize=9, fontname="helv"
                )
        for i, row in enumerate(lay["q"]):
            page.insert_text(
                (COLS_X[i // ROWS] - 2, row[0][1] + 3.5),
                f"{i + 1:>2}.",
                fontsize=10,
                fontname="helv",
            )
            for x, y in row:
                page.draw_circle((x, y), R, color=(0, 0, 0), width=0.8)
        page.insert_text(
            (70, H - 60),
            "NEXVARA javob varaqasi · burchakdagi qora kvadratlarni yopmang",
            fontsize=7,
            fontname="helv",
            color=(0.4, 0.4, 0.4),
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()
    return path


# ---------------- reading a photo ----------------
class Gray:
    """A grayscale image as bytes with (w, h)."""

    def __init__(self, data: bytes, w: int, h: int) -> None:
        self.data, self.w, self.h = data, w, h

    def at(self, x: float, y: float) -> int:
        xi, yi = int(x), int(y)
        if 0 <= xi < self.w and 0 <= yi < self.h:
            return self.data[yi * self.w + xi]
        return 255


def _gray(path: Path, long_side: int) -> Gray:
    import fitz

    pix = fitz.Pixmap(Path(path).read_bytes())
    if pix.n - pix.alpha != 1:
        pix = fitz.Pixmap(fitz.csGRAY, pix)
    if pix.alpha:
        pix = fitz.Pixmap(pix, 0)
    while max(pix.width, pix.height) > long_side * 1.9:
        pix.shrink(1)
    return (
        Gray(bytes(pix.samples), pix.width, pix.height)
        if pix.stride == pix.width
        else Gray(
            b"".join(
                pix.samples[r * pix.stride : r * pix.stride + pix.width] for r in range(pix.height)
            ),
            pix.width,
            pix.height,
        )
    )


def _threshold(img: Gray) -> int:
    """Ink vs paper (Otsu): works for a mostly white sheet and for a grey, dim photo alike."""
    hist = [0] * 256
    step = max(1, len(img.data) // 300_000)
    for v in img.data[::step]:
        hist[v] += 1
    total = sum(hist)
    sum_all = sum(i * c for i, c in enumerate(hist))
    best, t, w0, sum0 = -1.0, 128, 0, 0
    for v in range(256):
        w0 += hist[v]
        if w0 == 0 or w0 == total:
            continue
        sum0 += v * hist[v]
        m0, m1 = sum0 / w0, (sum_all - sum0) / (total - w0)
        between = w0 * (total - w0) * (m0 - m1) ** 2
        if between > best:
            best, t = between, v
    return max(40, min(t, 200))


def _blobs(img: Gray, t: int) -> list[dict[str, float]]:
    seen = bytearray(img.w * img.h)
    out = []
    d = img.data
    for start in range(img.w * img.h):
        if seen[start] or d[start] > t:
            continue
        stack = [start]
        seen[start] = 1
        xs = ys = area = 0
        x0, y0, x1, y1 = img.w, img.h, 0, 0
        while stack:
            i = stack.pop()
            y, x = divmod(i, img.w)
            area += 1
            xs += x
            ys += y
            x0, y0, x1, y1 = min(x0, x), min(y0, y), max(x1, x), max(y1, y)
            for j in (i - 1, i + 1, i - img.w, i + img.w):
                if 0 <= j < len(d) and not seen[j] and d[j] <= t and abs((j % img.w) - x) <= 1:
                    seen[j] = 1
                    stack.append(j)
        bw, bh = x1 - x0 + 1, y1 - y0 + 1
        out.append(
            {
                "x": xs / area,
                "y": ys / area,
                "area": area,
                "fill": area / (bw * bh),
                "aspect": bw / bh,
            }
        )
    return out


def _homography(src: list[tuple[float, float]], dst: list[tuple[float, float]]) -> list[float]:
    """3x3 (as 9 numbers, h33 = 1) mapping src points to dst points (Gaussian elimination)."""
    a, b = [], []
    for (x, y), (u, v) in zip(src, dst, strict=True):
        a.append([x, y, 1, 0, 0, 0, -u * x, -u * y])
        b.append(u)
        a.append([0, 0, 0, x, y, 1, -v * x, -v * y])
        b.append(v)
    n = 8
    m = [row[:] + [b[i]] for i, row in enumerate(a)]
    for c in range(n):
        p = max(range(c, n), key=lambda r: abs(m[r][c]))
        if abs(m[p][c]) < 1e-12:
            raise UserError("Varaqaning burchaklari noto‘g‘ri topildi")
        m[c], m[p] = m[p], m[c]
        for r in range(n):
            if r != c:
                f = m[r][c] / m[c][c]
                for k in range(c, n + 1):
                    m[r][k] -= f * m[c][k]
    return [m[i][n] / m[i][i] for i in range(n)] + [1.0]


def _map(h: list[float], x: float, y: float) -> tuple[float, float]:
    z = h[6] * x + h[7] * y + h[8]
    return (h[0] * x + h[1] * y + h[2]) / z, (h[3] * x + h[4] * y + h[5]) / z


def find_marks(img: Gray) -> list[tuple[float, float]]:
    t = min(_threshold(img), 120)
    side = max(img.w, img.h)
    lo, hi = (side * 0.012) ** 2, (side * 0.09) ** 2
    cand = [
        b
        for b in _blobs(img, t)
        if lo <= b["area"] <= hi and b["fill"] > 0.45 and 0.55 < b["aspect"] < 1.8
    ]
    if len(cand) < 4:
        raise UserError(
            "Varaqaning 4 ta qora burchagi topilmadi. Varaqni to‘liq, yorug‘ joyda, "
            "tepadan suratga oling."
        )
    # The corner squares are the biggest solid squares on the sheet; filled bubbles are smaller.
    cand.sort(key=lambda b: -b["area"])
    cand = [b for b in cand if b["area"] >= 0.5 * cand[3]["area"]]
    tl = min(cand, key=lambda b: b["x"] + b["y"])
    br = max(cand, key=lambda b: b["x"] + b["y"])
    tr = max(cand, key=lambda b: b["x"] - b["y"])
    bl = max(cand, key=lambda b: b["y"] - b["x"])
    pts = [(tl["x"], tl["y"]), (tr["x"], tr["y"]), (bl["x"], bl["y"]), (br["x"], br["y"])]
    if len(set(pts)) < 4:
        raise UserError("Varaq burchaklari aniqlanmadi, qayta suratga oling")
    return pts


def read_sheet(photo: Path, n: int, variants: int) -> dict[str, Any]:
    """Marked letters per question ("-" empty, "*" several) and the marked variant."""
    small = _gray(photo, 420)
    marks = find_marks(small)
    big = _gray(photo, 1400)
    sx, sy = big.w / small.w, big.h / small.h
    h = _homography(MARKS, [(x * sx, y * sy) for x, y in marks])
    t = _threshold(big)
    lay = layout(n, variants)

    def fill(x: float, y: float) -> float:
        cx, cy = _map(h, x, y)
        ex, _ = _map(h, x + R * 0.55, y)
        r = max(2.0, abs(ex - cx))
        dark = total = 0
        steps = 7
        for i in range(-steps, steps + 1):
            for j in range(-steps, steps + 1):
                if i * i + j * j > steps * steps:
                    continue
                total += 1
                if big.at(cx + i * r / steps, cy + j * r / steps) <= t:
                    dark += 1
        return dark / total

    def choose(cells: list[tuple[float, float]]) -> tuple[str, list[float]]:
        f = [fill(x, y) for x, y in cells]
        marked = [k for k, v in enumerate(f) if v >= 0.45]
        if not marked:
            return "-", f
        if len(marked) > 1:
            best = sorted(f, reverse=True)
            if best[0] - best[1] > 0.3:  # one clearly filled, another just smudged
                return str(f.index(best[0])), f
            return "*", f
        return str(marked[0]), f

    answers = []
    for row in lay["q"]:
        k, _f = choose(row)
        answers.append(LETTERS[int(k)] if k.isdigit() else k)
    v, _ = choose(lay["variant"]) if variants > 1 else ("0", [])
    return {"answers": "".join(answers), "variant": int(v) + 1 if v.isdigit() else None}


# ---------------- sets of sheets, keys, results ----------------
def _sets_file() -> Path:
    return folder() / "varaqalar.json"


def load_sets() -> list[dict[str, Any]]:
    try:
        data = json.loads(_sets_file().read_text(encoding="utf-8"))
        return data if isinstance(data, list) else []
    except (OSError, ValueError):
        return []


def _save_sets(sets: list[dict[str, Any]]) -> None:
    _sets_file().write_text(json.dumps(sets, ensure_ascii=False, indent=1), encoding="utf-8")


def clean_key(key: str, n: int) -> str:
    k = "".join(ch for ch in str(key).upper() if ch in LETTERS)
    if len(k) != n:
        raise UserError(f"Kalitda {n} ta javob bo‘lsin (A/B/C/D), hozir {len(k)} ta")
    return k


def create_set(
    title: str, n: int, variants: int, keys: dict[str, str] | None = None, copies: int = 1
) -> dict[str, Any]:
    title = (title or "Test").strip()[:70]
    layout(n, variants)
    sid = uuid.uuid4().hex[:8]
    keys = {str(v): clean_key(k, n) for v, k in (keys or {}).items() if str(k).strip()}
    pdf = make_sheet_pdf(folder() / f"Javob varaqasi {sid}.pdf", title, n, variants, copies)
    s = {
        "id": sid,
        "title": title,
        "questions": n,
        "variants": variants,
        "keys": keys,
        "pdf": str(pdf),
        "created": time.time(),
        "results": [],
    }
    sets = load_sets()
    sets.insert(0, s)
    _save_sets(sets[:50])
    return s


def get_set(sid: str) -> dict[str, Any]:
    for s in load_sets():
        if s["id"] == sid:
            return s
    raise LookupError("Javob varaqasi topilmadi")


def set_keys(sid: str, keys: dict[str, str]) -> dict[str, Any]:
    sets = load_sets()
    for s in sets:
        if s["id"] == sid:
            s["keys"] = {
                str(v): clean_key(k, s["questions"]) for v, k in keys.items() if str(k).strip()
            }
            _save_sets(sets)
            return s
    raise LookupError("Javob varaqasi topilmadi")


def mark(percent: float) -> int:
    """5-ball scale used in Uzbek schools."""
    return 5 if percent >= 86 else 4 if percent >= 71 else 3 if percent >= 56 else 2


def grade(sid: str, photo: Path, student: str = "") -> dict[str, Any]:
    s = get_set(sid)
    if not s["keys"]:
        raise UserError("Avval to‘g‘ri javoblar kalitini kiriting")
    got = read_sheet(photo, s["questions"], s["variants"])
    variant = got["variant"] or 1
    key = s["keys"].get(str(variant)) or next(iter(s["keys"].values()))
    right = sum(1 for a, b in zip(got["answers"], key, strict=True) if a == b)
    pct = round(right * 100.0 / s["questions"], 1)
    res = {
        "time": time.time(),
        "student": (student or "").strip()[:60] or f"O‘quvchi {len(s['results']) + 1}",
        "variant": variant,
        "answers": got["answers"],
        "key": key,
        "right": right,
        "total": s["questions"],
        "percent": pct,
        "mark": mark(pct),
        "wrong": [
            i + 1 for i, (a, b) in enumerate(zip(got["answers"], key, strict=True)) if a != b
        ],
    }
    sets = load_sets()
    for x in sets:
        if x["id"] == sid:
            x["results"].insert(0, res)
            x["results"] = x["results"][:500]
    _save_sets(sets)
    res["excel"] = str(write_excel(get_set(sid)))
    return res


def write_excel(s: dict[str, Any]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"
    ws.append([s["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Sana", "O‘quvchi", "Variant", "To‘g‘ri", "Jami", "Foiz", "Baho", "Xato savollar"])
    for c in ws[2]:
        c.font = Font(bold=True)
    for r in reversed(s["results"]):
        ws.append(
            [
                time.strftime("%d.%m.%Y %H:%M", time.localtime(r["time"])),
                r["student"],
                r["variant"],
                r["right"],
                r["total"],
                r["percent"],
                r["mark"],
                ", ".join(map(str, r["wrong"])),
            ]
        )
    for col, width in zip("ABCDEFGH", (17, 28, 9, 9, 7, 8, 7, 40), strict=True):
        ws.column_dimensions[col].width = width
    out = folder() / f"Natijalar {s['title'][:40]} {s['id']}.xlsx"
    wb.save(out)
    return out
